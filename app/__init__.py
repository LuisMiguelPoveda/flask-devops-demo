import os
import json
import time
import threading
import math
import re
import tempfile
import fcntl
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path
from io import BytesIO

import requests
from requests.exceptions import Timeout, RequestException
from markupsafe import Markup, escape
from werkzeug.exceptions import RequestEntityTooLarge
from werkzeug.utils import secure_filename
from PyPDF2 import PdfReader
from pptx import Presentation
from sqlalchemy import event
from sqlalchemy.engine import Engine
from sqlalchemy.exc import OperationalError

from flask import Flask, render_template, request, redirect, url_for, flash, jsonify, session, send_file
from flask_wtf.csrf import CSRFProtect
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address
from flask_migrate import Migrate
from flask_login import (
    LoginManager,
    login_user,
    login_required,
    logout_user,
    current_user,
)
from werkzeug.security import check_password_hash, generate_password_hash

from .models import (
    db,
    User,
    Subject,
    Note,
    FlashcardDeck,
    Job,
    AskProfeMessage,
    ProfeSessionLock,
    StudentProfile,
    SubjectExam,
    NoteSourceFile,
    TaskItem,
)


ALLOWED_EXTENSIONS = {"txt", "pdf", "pptx"}
# Permitimos hasta ~50 MB por archivo
MAX_UPLOAD_BYTES = 50 * 1024 * 1024
LLM_PROFILE_CHOICES = [
    {
        "id": "vram_4gb",
        "label": "4 GB (conservador)",
        "chunk_tokens": 1500,
        "chunk_overlap": 250,
        "summary_max_tokens": 1200,
        "summary_min_tokens": 120,
        "summary_ratio": 0.35,
    },
    {
        "id": "vram_8gb",
        "label": "8 GB (equilibrado)",
        "chunk_tokens": 2200,
        "chunk_overlap": 350,
        "summary_max_tokens": 2200,
        "summary_min_tokens": 120,
        "summary_ratio": 0.45,
    },
    {
        "id": "vram_16gb",
        "label": "16 GB (alto)",
        "chunk_tokens": 3000,
        "chunk_overlap": 500,
        "summary_max_tokens": 3500,
        "summary_min_tokens": 120,
        "summary_ratio": 0.5,
    },
]
LLM_PROFILE_PRESETS = {choice["id"]: choice for choice in LLM_PROFILE_CHOICES}
DEFAULT_LLM_PROFILE = "vram_4gb"
DEFAULT_LLM_LIMITS = LLM_PROFILE_PRESETS[DEFAULT_LLM_PROFILE]
MAX_SUMMARY_TOKENS = DEFAULT_LLM_LIMITS["summary_max_tokens"]
SUMMARY_MIN_TOKENS = DEFAULT_LLM_LIMITS["summary_min_tokens"]
SUMMARY_TOKEN_RATIO = DEFAULT_LLM_LIMITS["summary_ratio"]
DEFAULT_CHUNK_TOKENS = DEFAULT_LLM_LIMITS["chunk_tokens"]
DEFAULT_CHUNK_OVERLAP = DEFAULT_LLM_LIMITS["chunk_overlap"]
FLASHCARD_CHUNK_COUNTS = (5, 10, 15, 20)
FLASHCARD_CHUNK_DEFAULT = FLASHCARD_CHUNK_COUNTS[0]
TASK_WINDOW_OPTIONS = [
    {"value": 0, "label": "Todos"},
    {"value": 7, "label": "1 semana"},
    {"value": 14, "label": "2 semanas"},
    {"value": 30, "label": "1 mes"},
]
TASK_WINDOW_VALUES = {opt["value"] for opt in TASK_WINDOW_OPTIONS}
TASK_WINDOW_LABELS = {opt["value"]: opt["label"] for opt in TASK_WINDOW_OPTIONS}
TASK_WINDOW_DEFAULT = 14
LLM_OFFLINE_LABEL = "Motor LLM no encontrado"
HEX_COLOR_RE = re.compile(r"^#[0-9a-fA-F]{6}$")
# evita arrancar dos workers en procesos reloader
_worker_started = False
_worker_lock_handle = None

limiter = Limiter(key_func=get_remote_address, default_limits=[])
migrate = Migrate()
JOB_RETRY_SECONDS = 30
INCOMPLETE_JOB_STATUSES = ("pending", "running")


def chunk_text_with_overlap(
    text: str,
    max_tokens: int = DEFAULT_CHUNK_TOKENS,
    overlap: int = DEFAULT_CHUNK_OVERLAP,
) -> list[str]:
    """
    Corta el texto en fragmentos aproximados de tokens (palabras) con solapamiento
    para minimizar la pérdida de contexto entre partes.
    """
    tokens = (text or "").split()
    if not tokens:
        return []

    max_tokens = max(1, max_tokens)
    overlap = max(0, min(overlap, max_tokens - 1))
    step = max_tokens - overlap if max_tokens > overlap else 1

    chunks: list[str] = []
    start = 0
    while start < len(tokens):
        end = min(len(tokens), start + max_tokens)
        chunk_tokens = tokens[start:end]
        chunks.append(" ".join(chunk_tokens))
        if end >= len(tokens):
            break
        start += step
    return chunks


def estimate_summary_max_tokens(
    text: str,
    max_tokens: int = MAX_SUMMARY_TOKENS,
    min_tokens: int = SUMMARY_MIN_TOKENS,
    ratio: float = SUMMARY_TOKEN_RATIO,
) -> int:
    word_count = len((text or "").split())
    target = int(word_count * ratio)
    if target <= 0:
        target = min_tokens
    return max(min_tokens, min(max_tokens, target))


def normalize_subject_color(value: str | None) -> str | None:
    if not value:
        return None
    value = value.strip()
    if not value:
        return None
    return value if HEX_COLOR_RE.match(value) else None


def resolve_llm_profile(user_id: int | None) -> str:
    if not user_id:
        return DEFAULT_LLM_PROFILE
    profile = StudentProfile.query.filter_by(user_id=user_id).first()
    if profile and profile.llm_profile:
        candidate = profile.llm_profile.strip()
        if candidate in LLM_PROFILE_PRESETS:
            return candidate
    return DEFAULT_LLM_PROFILE


def get_llm_limits(user_id: int | None) -> dict:
    profile_key = resolve_llm_profile(user_id)
    return LLM_PROFILE_PRESETS.get(profile_key, DEFAULT_LLM_LIMITS)


def resolve_task_window_days(profile: StudentProfile | None) -> int:
    if profile and profile.task_window_days in TASK_WINDOW_VALUES:
        return profile.task_window_days
    return TASK_WINDOW_DEFAULT


def calendar_window_range(profile: StudentProfile | None):
    window_days = resolve_task_window_days(profile)
    today = datetime.utcnow().date()
    if window_days == 0:
        return window_days, None, None
    end_date = today + timedelta(days=max(1, window_days) - 1)
    return window_days, today, end_date


def build_calendar_items(user_id: int, start_date, end_date, today=None) -> list[dict]:
    if today is None:
        today = datetime.utcnow().date()
    tasks_query = TaskItem.query.filter_by(user_id=user_id)
    exams_query = SubjectExam.query.join(Subject, SubjectExam.subject_id == Subject.id).filter(
        Subject.user_id == user_id
    )
    if start_date is not None and end_date is not None:
        tasks_query = tasks_query.filter(TaskItem.due_date >= start_date, TaskItem.due_date <= end_date)
        exams_query = exams_query.filter(SubjectExam.exam_date >= start_date, SubjectExam.exam_date <= end_date)
    tasks = tasks_query.order_by(TaskItem.due_date.asc(), TaskItem.title.asc()).all()
    exams = exams_query.order_by(SubjectExam.exam_date.asc(), SubjectExam.tema.asc()).all()

    def relative_label(target_date):
        delta_days = (target_date - today).days
        if delta_days < 0:
            days = abs(delta_days)
            return "Hace 1 día" if days == 1 else f"Hace {days} días"
        if delta_days == 0:
            return "Hoy"
        if delta_days == 1:
            return "Mañana"
        return f"En {delta_days} días"

    items: list[dict] = []
    for task in tasks:
        items.append(
            {
                "kind": "task",
                "date": task.due_date,
                "title": task.title,
                "notes": task.notes,
                "subject": task.subject,
                "task": task,
                "relative_label": relative_label(task.due_date),
            }
        )
    for exam in exams:
        items.append(
            {
                "kind": "exam",
                "date": exam.exam_date,
                "title": exam.tema,
                "subject": exam.subject,
                "exam": exam,
                "relative_label": relative_label(exam.exam_date),
            }
        )

    items.sort(key=lambda item: (item["date"], item["kind"], (item["title"] or "").lower()))
    return items


def payload_int(payload: dict | None, key: str) -> int | None:
    if not isinstance(payload, dict):
        return None
    value = payload.get(key)
    if value is None:
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def has_incomplete_jobs_for_payload(user_id: int, key: str, value: int | None) -> bool:
    if value is None:
        return False
    jobs = Job.query.filter(
        Job.user_id == user_id,
        Job.status.in_(INCOMPLETE_JOB_STATUSES),
    ).all()
    for job in jobs:
        if payload_int(job.payload or {}, key) == value:
            return True
    return False


def job_is_cancelled(job_id: int | None) -> bool:
    if not job_id:
        return False
    status = db.session.query(Job.status).filter_by(id=job_id).scalar()
    return status == "cancelled"


def commit_if_not_cancelled(job_id: int) -> bool:
    with db.session.no_autoflush:
        if job_is_cancelled(job_id):
            db.session.rollback()
            return False
    db.session.commit()
    return True


def build_queue_groups(user_id: int) -> list[dict]:
    jobs = (
        Job.query.filter(
            Job.user_id == user_id,
            Job.type.in_(
                (
                    "file_import",
                    "note_ai",
                    "note_ai_chunk",
                    "flashcards_ai_new",
                    "flashcards_ai_append",
                    "flashcards_ai_chunk",
                )
            ),
        )
        .order_by(Job.created_at.asc(), Job.id.asc())
        .all()
    )
    notes = {note.id: note for note in Note.query.filter_by(user_id=user_id).all()}
    decks = {deck.id: deck for deck in FlashcardDeck.query.filter_by(user_id=user_id).all()}

    note_groups: dict[int, dict] = {}
    file_groups: list[dict] = []

    for job in jobs:
        payload = job.payload or {}
        if job.type == "file_import":
            if job.status in INCOMPLETE_JOB_STATUSES:
                file_groups.append(
                    {
                        "kind": "file_import",
                        "job_id": job.id,
                        "title": payload.get("filename") or "Archivo en cola",
                        "status": job.status,
                        "created_at": job.created_at,
                    }
                )
            continue

        note_id = payload_int(payload, "note_id")
        if not note_id:
            continue
        group = note_groups.setdefault(
            note_id,
            {
                "note_id": note_id,
                "jobs": [],
                "earliest_incomplete": None,
            },
        )
        group["jobs"].append(job)
        if job.status in INCOMPLETE_JOB_STATUSES:
            job_created = job.created_at or datetime.min
            if not group["earliest_incomplete"] or job_created < group["earliest_incomplete"]:
                group["earliest_incomplete"] = job_created

    groups: list[dict] = []
    for note_id, group in note_groups.items():
        start_at = group.get("earliest_incomplete")
        if not start_at:
            continue
        jobs = [
            job
            for job in group["jobs"]
            if (job.created_at or datetime.min) >= start_at
        ]
        if not any(job.status in INCOMPLETE_JOB_STATUSES for job in jobs):
            continue
        deck_ids = set()
        for job in jobs:
            deck_id = payload_int(job.payload or {}, "deck_id")
            if deck_id:
                deck_ids.add(deck_id)
        summary_jobs = [job for job in jobs if job.type in ("note_ai", "note_ai_chunk")]
        flash_jobs = [
            job
            for job in jobs
            if job.type in ("flashcards_ai_new", "flashcards_ai_append", "flashcards_ai_chunk")
        ]
        summary_done = sum(1 for job in summary_jobs if job.status == "success")
        flash_done = sum(1 for job in flash_jobs if job.status == "success")
        deck_titles = [decks[d].title for d in deck_ids if d in decks and decks[d].title]
        note = notes.get(note_id)
        created_at = min(
            (job.created_at for job in jobs if job.created_at),
            default=start_at,
        )
        groups.append(
            {
                "kind": "note",
                "note_id": note_id,
                "title": note.title if note else f"Apunte #{note_id}",
                "summary_done": summary_done,
                "summary_total": len(summary_jobs),
                "flash_done": flash_done,
                "flash_total": len(flash_jobs),
                "deck_titles": deck_titles,
                "created_at": created_at,
            }
        )

    groups.extend(file_groups)
    groups.sort(key=lambda item: item.get("created_at") or datetime.min, reverse=True)
    return groups


def build_note_chunks_map(
    user_id: int,
    notes: list[Note],
    max_tokens: int = DEFAULT_CHUNK_TOKENS,
    overlap: int = DEFAULT_CHUNK_OVERLAP,
) -> dict[int, dict]:
    """
    Devuelve un dict {note_id: {"chunks": [str], "total": int}}.
    Usa los trabajos note_ai_chunk si existen para respetar el número de fragmentos originales;
    si no hay trabajos, divide el contenido del apunte.
    """
    if not notes:
        return {}

    note_ids = {n.id for n in notes if n.id}
    jobs = Job.query.filter(Job.user_id == user_id, Job.type == "note_ai_chunk").all()
    by_note: dict[int, list[tuple[int, str, int]]] = {}
    for job in jobs:
        payload = job.payload or {}
        note_id = payload.get("note_id")
        if note_id not in note_ids:
            continue
        try:
            idx = max(0, int(payload.get("chunk_index") or 0))
        except (ValueError, TypeError):
            idx = 0
        try:
            total = max(1, int(payload.get("total_chunks") or 1))
        except (ValueError, TypeError):
            total = 1
        text = payload.get("text") or ""
        by_note.setdefault(note_id, []).append((idx, text, total))

    result: dict[int, dict] = {}
    for note in notes:
        if note.id in by_note:
            chunks_info = sorted(by_note[note.id], key=lambda t: t[0])
            chunks = [t[1] for t in chunks_info if t[1]]
            total_max = max((t[2] for t in chunks_info if t[2]), default=len(chunks))
            result[note.id] = {"chunks": chunks, "total": total_max or len(chunks) or 1}
            continue

        chunks = chunk_text_with_overlap(note.content or "", max_tokens=max_tokens, overlap=overlap)
        result[note.id] = {"chunks": chunks or [note.content or ""], "total": len(chunks) or 1}

    return result


def simple_format_note(text: str) -> Markup:
    """
    Convierte texto plano con encabezados '#' y viñetas '*'/'-' en HTML simple.
    Soporta **negrita** y *cursiva*. Escapa contenido para evitar XSS.
    """
    import re
    from html import unescape as html_unescape

    lines = html_unescape(text or "").splitlines()
    html_parts = []
    in_list = False

    def close_list():
        nonlocal in_list
        if in_list:
            html_parts.append("</ul>")
            in_list = False

    for line in lines:
        stripped = line.strip()
        if not stripped:
            close_list()
            continue

        def fmt_inline(txt: str) -> str:
            esc = escape(txt)
            esc = re.sub(r"\*\*(.+?)\*\*", lambda m: f"<strong>{escape(m.group(1))}</strong>", esc)
            esc = re.sub(r"\*(.+?)\*", lambda m: f"<em>{escape(m.group(1))}</em>", esc)
            return esc

        if stripped == "---":
            close_list()
            html_parts.append("<hr />")
        elif stripped.startswith("###"):
            close_list()
            html_parts.append(f"<h4>{fmt_inline(stripped.lstrip('#').strip())}</h4>")
        elif stripped.startswith("##"):
            close_list()
            html_parts.append(f"<h3>{fmt_inline(stripped.lstrip('#').strip())}</h3>")
        elif stripped.startswith("#"):
            close_list()
            html_parts.append(f"<h2>{fmt_inline(stripped.lstrip('#').strip())}</h2>")
        elif stripped.startswith(("* ", "- ")):
            if not in_list:
                html_parts.append("<ul>")
                in_list = True
            html_parts.append(f"<li>{fmt_inline(stripped[2:].strip())}</li>")
        else:
            close_list()
            html_parts.append(f"<p>{fmt_inline(stripped)}</p>")

    close_list()
    return Markup("".join(html_parts))


def fetch_ask_profe_history(user_id: int, limit: int = 12) -> list[dict]:
    rows = (
        AskProfeMessage.query.filter_by(user_id=user_id)
        .order_by(AskProfeMessage.created_at.desc(), AskProfeMessage.id.desc())
        .limit(limit)
        .all()
    )
    rows.reverse()
    return [{"role": row.role, "content": row.content} for row in rows]


def prune_ask_profe_history(user_id: int, keep: int = 12) -> None:
    extras = (
        AskProfeMessage.query.filter_by(user_id=user_id)
        .order_by(AskProfeMessage.created_at.desc(), AskProfeMessage.id.desc())
        .offset(keep)
        .all()
    )
    for msg in extras:
        db.session.delete(msg)


def fetch_tema_options(user_id: int, subject_id: int | None = None) -> list[str]:
    q = (
        db.session.query(SubjectExam.tema)
        .join(Subject, SubjectExam.subject_id == Subject.id)
        .filter(Subject.user_id == user_id)
    )
    if subject_id:
        q = q.filter(SubjectExam.subject_id == subject_id)
    rows = q.distinct().order_by(SubjectExam.tema.asc()).all()
    return [r[0] for r in rows if r[0]]


def resolve_default_model(app, user_id: int | None, available_models: list[str] | None = None) -> str:
    if available_models and len(available_models) == 1 and available_models[0] == LLM_OFFLINE_LABEL:
        return LLM_OFFLINE_LABEL
    model = app.config["LMSTUDIO_MODEL"]
    if user_id:
        profile = StudentProfile.query.filter_by(user_id=user_id).first()
        if profile and profile.default_model:
            model = profile.default_model
    if available_models and model not in available_models:
        return available_models[0] if available_models else model
    return model


def is_setup_complete(user_id: int) -> bool:
    if not user_id:
        return False
    profile = StudentProfile.query.filter_by(user_id=user_id).first()
    return profile is not None


def allowed_file(filename: str) -> bool:
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def save_upload_stream(upload, dest_path: str, max_bytes: int) -> tuple[int, str | None]:
    total = 0
    try:
        with open(dest_path, "wb") as handle:
            while True:
                chunk = upload.stream.read(1024 * 1024)
                if not chunk:
                    break
                total += len(chunk)
                if total > max_bytes:
                    break
                handle.write(chunk)
    except Exception:
        if os.path.exists(dest_path):
            os.remove(dest_path)
        raise

    if total == 0:
        if os.path.exists(dest_path):
            os.remove(dest_path)
        return 0, "empty"
    if total > max_bytes:
        if os.path.exists(dest_path):
            os.remove(dest_path)
        return total, "too_large"
    return total, None


def extract_pdf_text(file_bytes: bytes) -> str:
    """Extrae texto simple desde PDF usando PyPDF2."""
    reader = PdfReader(BytesIO(file_bytes))
    chunks: list[str] = []
    for page in reader.pages:
        try:
            txt = page.extract_text() or ""
        except Exception:
            txt = ""
        if txt:
            chunks.append(txt)
    return "\n".join(chunks)


def extract_pptx_text(file_bytes: bytes) -> str:
    """Extrae texto simple desde PPTX usando python-pptx."""
    prs = Presentation(BytesIO(file_bytes))
    chunks: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text or ""
                if text.strip():
                    chunks.append(text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    row_text = " ".join(cell.text.strip() for cell in row.cells if cell.text)
                    if row_text:
                        chunks.append(row_text)
    return "\n".join(chunks)


def fetch_models(app):
    api_base = app.config["LMSTUDIO_API_BASE"].rstrip("/")
    try:
        resp = requests.get(f"{api_base}/models", timeout=5)
        resp.raise_for_status()
        data = resp.json()
        all_ids = [m["id"] for m in data.get("data", [])]
        models = [mid for mid in all_ids if not mid.startswith("text-embedding-")]
        return models or [LLM_OFFLINE_LABEL]
    except Exception:
        return [LLM_OFFLINE_LABEL]


def lmstudio_chat(
    app,
    model: str,
    messages: list[dict],
    response_format: dict | None = None,
    max_tokens: int | None = None,
    temperature: float = 0.4,
) -> str:
    api_base = app.config["LMSTUDIO_API_BASE"].rstrip("/")
    timeout_s = app.config["LMSTUDIO_TIMEOUT"]
    payload = {"model": model, "messages": messages, "temperature": temperature}
    if response_format:
        payload["response_format"] = response_format
    if max_tokens is not None:
        payload["max_tokens"] = max_tokens
    resp = requests.post(f"{api_base}/chat/completions", json=payload, timeout=timeout_s)
    resp.raise_for_status()
    data = resp.json()
    return data["choices"][0]["message"]["content"].strip()


def lmstudio_summarize_text(
    app,
    model: str,
    subject: str,
    title: str,
    exam_date: str,
    filename: str,
    text: str,
    chunk_index: int | None = None,
    total_chunks: int | None = None,
    user_id: int | None = None,
) -> str:
    system_prompt = (
        "Actúa como un especialista en pedagogía y diseño instruccional con experiencia en TDAH y carga cognitiva.\n\n"
        "Tu tarea es transformar el TEXTO FUENTE en APUNTES GUIADOS altamente estructurados,\n"
        "optimizados para estudiantes con déficit de atención.\n\n"
        "REGLAS GENERALES (obligatorias):\n"
        "1. Mantén SIEMPRE la misma macroestructura y los mismos encabezados.\n"
        "2. Usa fragmentación cognitiva (chunking): bloques cortos y visualmente claros.\n"
        "3. Reduce texto continuo: prioriza bullets jerárquicos, tablas y esquemas.\n"
        "4. Señaliza explícitamente lo importante (clave, confusión común, examinable).\n"
        "5. No añadas información nueva: solo reorganiza y clarifica el texto dado.\n"
        "6. Lenguaje claro, directo, sin metáforas ni digresiones.\n"
        "7. Cada bloque debe poder leerse de forma independiente en <30 segundos.\n"
        "8. Máximo 12–15 líneas por sección principal.\n"
        "9. Usa siempre palabras clave antes de explicaciones breves.\n"
        "10. Evita listas planas: mínimo dos niveles de jerarquía cuando haya listas.\n\n"
        "FORMATO FIJO (NO MODIFICAR):\n\n"
        "────────────────────────\n"
        " TEMA:\n"
        "[Nombre claro y conciso del tema]\n\n"
        " IDEA CENTRAL (1 frase):\n"
        "• [Qué es / para qué sirve]\n\n"
        "────────────────────────\n"
        " CONCEPTOS CLAVE:\n"
        "• Concepto 1\n"
        "  – Definición corta\n"
        "  – Ejemplo mínimo (si aplica)\n"
        "• Concepto 2\n"
        "  – Definición corta\n"
        "  – Ejemplo mínimo\n\n"
        "────────────────────────\n"
        " RELACIONES IMPORTANTES:\n"
        "• [Concepto A] → [Concepto B]\n"
        "  – Tipo de relación (causa, consecuencia, contraste, parte–todo)\n\n"
        "────────────────────────\n"
        " PASOS / PROCESO / ESTRUCTURA (si aplica):\n"
        "1. Paso 1 — palabra clave\n"
        "   – Qué ocurre\n"
        "2. Paso 2 — palabra clave\n"
        "   – Qué ocurre\n\n"
        "────────────────────────\n"
        " ERRORES O CONFUSIONES COMUNES:\n"
        "• Error frecuente\n"
        "  – Por qué es incorrecto\n\n"
        "────────────────────────\n"
        " LO QUE SÍ ENTRA EN EXAMEN / EVALUACIÓN:\n"
        "• Hecho, definición o relación clave\n\n"
        "────────────────────────\n"
        " HUECOS PARA COMPLETAR (guided notes):\n"
        "• _______________________________\n"
        "• _______________________________\n\n"
        "────────────────────────\n"
        " CONEXIÓN CON OTROS APUNTES:\n"
        "• Se relaciona con: [tema previo / tema siguiente]\n"
        "• Idea puente: ____________________\n\n"
        "────────────────────────\n\n"
        "ENTRADA:\n"
        "[TEXTO FUENTE AQUÍ]\n\n"
        "SALIDA:\n"
        "Apuntes siguiendo EXACTAMENTE el formato indicado."
    )
    chunk_meta_line = ""
    if total_chunks and total_chunks > 1:
        human_idx = (chunk_index or 0) + 1
        chunk_meta_line = f"Parte: {human_idx}/{total_chunks} (solo para contexto; no lo menciones en tu respuesta)."
    user_prompt = (
        f"Asignatura: {subject}\n"
        f"Título: {title}\n"
        f"Fecha de examen: {exam_date}\n"
        f"Archivo: {filename}\n"
        f"{chunk_meta_line}\n\n"
        f"TEXTO A RESUMIR:\n{text}"
    )
    limits = get_llm_limits(user_id)
    max_tokens = estimate_summary_max_tokens(
        text,
        max_tokens=limits["summary_max_tokens"],
        min_tokens=limits["summary_min_tokens"],
        ratio=limits["summary_ratio"],
    )
    return lmstudio_chat(
        app,
        model,
        [{"role": "system", "content": system_prompt}, {"role": "user", "content": user_prompt}],
        max_tokens=max_tokens,
    )


def _parse_and_validate_flashcards_json(raw: str, expected_count: int) -> list[dict]:
    """
    Esperamos EXACTAMENTE expected_count flashcards:
    [
      {"question": str, "options": [str,str,str,str], "correct_index": 0..3}
    ]
    """
    # A veces el modelo devuelve ```json ... ```
    cleaned = raw.strip()
    if cleaned.startswith("```"):
        cleaned = cleaned.strip("`")
        # si trae 'json\n'
        cleaned = cleaned.replace("json\n", "", 1).strip()

    data = json.loads(cleaned)

    if not isinstance(data, list) or len(data) != expected_count:
        raise ValueError(f"El JSON debe ser una lista de {expected_count} flashcards.")

    for i, card in enumerate(data, start=1):
        if not isinstance(card, dict):
            raise ValueError(f"Flashcard {i} no es un objeto JSON.")
        q = card.get("question")
        opts = card.get("options")
        idx = card.get("correct_index")
        if not isinstance(q, str) or not q.strip():
            raise ValueError(f"Flashcard {i} tiene 'question' inválida.")
        if not isinstance(opts, list) or len(opts) != 4 or not all(isinstance(o, str) and o.strip() for o in opts):
            raise ValueError(f"Flashcard {i} debe tener 4 'options' (strings).")
        if not isinstance(idx, int) or idx < 0 or idx > 3:
            raise ValueError(f"Flashcard {i} debe tener 'correct_index' entre 0 y 3.")
    return data


def lmstudio_generate_flashcards(app, model: str, note: Note, count: int = 5, source_text: str | None = None) -> list[dict]:
    """
    Genera N flashcards a partir del contenido del apunte/resumen.
    Devuelve lista de dicts validada.
    """
    count = max(1, min(count, 50))
    system_prompt = (
        f"Genera exactamente {count} flashcards de examen a partir del texto. "
        "Devuelve SOLO un JSON válido (sin texto extra, sin markdown). "
        "Formato: "
        "["
        '{"question":"...","options":["A","B","C","D"],"correct_index":0},'
        "..."
        "]. "
        "Las preguntas deben ser autocontenidas: no uses referencias como "
        "\"según el texto\", \"en el fragmento\", \"¿qué se mencionó?\" o similares."
    )

    text = source_text if source_text is not None else (note.content or "")
    user_prompt = (
        f"Asignatura: {note.subject.name}\n"
        f"Título: {note.title}\n"
        f"Fecha de examen: {note.exam_date.isoformat() if note.exam_date else 'No indicada'}\n\n"
        f"TEXTO:\n{text}\n\n"
        "Crea preguntas potenciales de examen que se entiendan sin contexto adicional, "
        "4 opciones, solo 1 correcta. No uses referencias al texto o a un fragmento, "
        "y no añadas contexto externo."
    )

    schema = {
        "type": "array",
        "minItems": count,
        "maxItems": count,
        "items": {
            "type": "object",
            "properties": {
                "question": {"type": "string", "minLength": 1},
                "options": {
                    "type": "array",
                    "minItems": 4,
                    "maxItems": 4,
                    "items": {"type": "string", "minLength": 1},
                },
                "correct_index": {"type": "integer", "minimum": 0, "maximum": 3},
            },
            "required": ["question", "options", "correct_index"],
            "additionalProperties": False,
        },
    }

    raw = lmstudio_chat(
        app,
        model,
        [{"role": "system", "content": system_prompt}, {"role": "user", "content": user_prompt}],
        response_format={
            "type": "json_schema",
            "json_schema": {"name": "flashcards", "schema": schema},
        },
        temperature=0.2,
    )
    cards = _parse_and_validate_flashcards_json(raw, expected_count=count)
    if len(cards) != count:
        raise ValueError(f"El modelo devolvió {len(cards)} flashcards, se esperaban {count}.")
    return cards


def process_job(app, job: Job):
    """
    Ejecuta un trabajo de la cola. Devuelve (status, message, error_message).
    """
    with app.app_context():
        try:
            if job.type == "file_import":
                payload = job.payload or {}
                user_id = payload.get("user_id")
                subject_id = payload.get("subject_id")
                exam_id = payload.get("exam_id")
                filename = payload.get("filename") or "input.txt"
                file_path = payload.get("file_path") or ""
                content_type = payload.get("content_type") or ""
                model = payload.get("model") or app.config["LMSTUDIO_MODEL"]

                if not file_path or not os.path.exists(file_path):
                    return "error", None, "Archivo subido no encontrado."

                subject = Subject.query.filter_by(id=subject_id, user_id=user_id).first()
                if not subject:
                    return "error", None, "Asignatura inválida para el archivo."
                exam = SubjectExam.query.filter_by(id=exam_id, subject_id=subject.id).first()
                if not exam:
                    return "error", None, "Examen inválido para el archivo."

                file_size = os.path.getsize(file_path)
                if file_size == 0:
                    return "error", None, "El archivo está vacío."
                if file_size > MAX_UPLOAD_BYTES:
                    return "error", None, f"Archivo demasiado grande. Máximo {MAX_UPLOAD_BYTES // 1024} KB."

                with open(file_path, "rb") as handle:
                    file_bytes = handle.read()
                if not file_bytes:
                    return "error", None, "El archivo está vacío."

                ext = filename.rsplit(".", 1)[1].lower() if "." in filename else ""
                if ext == "pdf":
                    file_mime = content_type or "application/pdf"
                    file_text = extract_pdf_text(file_bytes)
                elif ext == "pptx":
                    file_mime = (
                        content_type
                        or "application/vnd.openxmlformats-officedocument.presentationml.presentation"
                    )
                    file_text = extract_pptx_text(file_bytes)
                else:
                    file_mime = content_type or "text/plain"
                    try:
                        file_text = file_bytes.decode("utf-8")
                    except UnicodeDecodeError:
                        file_text = file_bytes.decode("utf-8", errors="ignore")
                file_text = (file_text or "").strip()
                if not file_text:
                    return "error", None, "El archivo no contiene texto legible."
                limits = get_llm_limits(user_id)
                chunks = chunk_text_with_overlap(
                    file_text,
                    max_tokens=limits["chunk_tokens"],
                    overlap=limits["chunk_overlap"],
                )
                if not chunks:
                    return "error", None, "No se pudo dividir el texto para IA."
                if job_is_cancelled(job.id):
                    if file_path and os.path.exists(file_path):
                        os.remove(file_path)
                    return "cancelled", None, None

                exam_date_str = exam.exam_date.isoformat() if exam.exam_date else "No indicada"
                base_name = Path(filename).stem if filename else subject.name
                note_title = f"{exam.tema} - {base_name}".strip(" -")
                if exam_date_str != "No indicada":
                    note_title = f"{note_title} ({exam_date_str})"
                deck_title = exam.tema

                note = Note(
                    user_id=user_id,
                    subject_id=subject.id,
                    title=note_title,
                    exam_date=exam.exam_date,
                    original_filename=filename,
                    content=f"{note_title}\n\n",
                    ai_used=True,
                )
                db.session.add(note)
                db.session.flush()
                db.session.add(
                    NoteSourceFile(
                        note_id=note.id,
                        filename=filename or "input.txt",
                        content_type=file_mime or "application/octet-stream",
                        data=file_bytes,
                    )
                )

                deck = FlashcardDeck.query.filter_by(
                    user_id=user_id,
                    subject_id=subject.id,
                    exam_date=exam.exam_date,
                    title=deck_title,
                ).first()
                if not deck:
                    deck = FlashcardDeck(
                        user_id=user_id,
                        subject_id=subject.id,
                        title=deck_title,
                        exam_date=exam.exam_date,
                        source_note_id=note.id,
                        flashcards=[],
                    )
                    db.session.add(deck)
                    db.session.flush()

                deck_size_before = len(deck.flashcards or [])
                for idx, chunk in enumerate(chunks):
                    db.session.add(
                        Job(
                            user_id=user_id,
                            type="note_ai_chunk",
                            payload={
                                "user_id": user_id,
                                "subject_id": subject.id,
                                "note_id": note.id,
                                "title": note_title,
                                "exam_date": exam_date_str,
                                "filename": filename or "input.txt",
                                "text": chunk,
                                "model": model,
                                "chunk_index": idx,
                                "total_chunks": len(chunks),
                            },
                        )
                    )
                    db.session.add(
                        Job(
                            user_id=user_id,
                            type="flashcards_ai_chunk",
                            payload={
                                "user_id": user_id,
                                "note_id": note.id,
                                "deck_id": deck.id,
                                "deck_size_before": deck_size_before,
                                "model": model,
                                "count": FLASHCARD_CHUNK_DEFAULT,
                                "chunk_index": idx,
                                "total_chunks": len(chunks),
                                "text": chunk,
                            },
                        )
                    )
                if not commit_if_not_cancelled(job.id):
                    if file_path and os.path.exists(file_path):
                        os.remove(file_path)
                    return "cancelled", None, None
                if file_path and os.path.exists(file_path):
                    os.remove(file_path)
                return "success", f"Archivo procesado: {filename}", None

            if job.type == "note_ai":
                payload = job.payload or {}
                user_id = payload.get("user_id")
                subject_id = payload.get("subject_id")
                title = payload.get("title")
                exam_date_str = payload.get("exam_date")
                filename = payload.get("filename") or "input.txt"
                text = payload.get("text") or ""
                model = payload.get("model") or app.config["LMSTUDIO_MODEL"]

                subject = Subject.query.filter_by(id=subject_id, user_id=user_id).first()
                user = User.query.get(user_id)
                if not subject or not user:
                    return "error", None, "Asignatura o usuario inválido."

                content = lmstudio_summarize_text(
                    app,
                    model,
                    subject.name,
                    title,
                    exam_date_str,
                    filename,
                    text,
                    user_id=user_id,
                )
                if not content.strip():
                    return "error", None, "El modelo devolvió un resumen vacío."
                # Prepend title to content to ensure first line carries it.
                content = f"{title}\n\n{content}"

                exam_date = datetime.strptime(exam_date_str, "%Y-%m-%d").date() if exam_date_str and exam_date_str != "No indicada" else None
                note = Note(
                    user_id=user.id,
                    subject_id=subject.id,
                    title=title,
                    exam_date=exam_date,
                    original_filename=filename,
                    content=content,
                    ai_used=True,
                )
                db.session.add(note)
                if not commit_if_not_cancelled(job.id):
                    return "cancelled", None, None
                return "success", f"Resumen listo: {title}", None

            if job.type == "note_ai_chunk":
                payload = job.payload or {}
                user_id = payload.get("user_id")
                subject_id = payload.get("subject_id")
                note_id = payload.get("note_id")
                filename = payload.get("filename") or "input.txt"
                text = payload.get("text") or ""
                model = payload.get("model") or app.config["LMSTUDIO_MODEL"]
                chunk_index = max(0, int(payload.get("chunk_index") or 0))
                total_chunks = max(1, int(payload.get("total_chunks") or 1))
                exam_date_str = payload.get("exam_date") or "No indicada"

                note = Note.query.filter_by(id=note_id, user_id=user_id).first()
                subject_lookup_id = subject_id or (note.subject_id if note else None)
                subject = Subject.query.filter_by(id=subject_lookup_id, user_id=user_id).first() if subject_lookup_id else None
                if not note or not subject:
                    return "error", None, "Apunte o asignatura inválidos."

                exam_date_str = exam_date_str or (note.exam_date.isoformat() if note.exam_date else "No indicada")
                summary = lmstudio_summarize_text(
                    app,
                    model,
                    subject.name,
                    note.title,
                    exam_date_str,
                    filename,
                    text,
                    chunk_index=chunk_index,
                    total_chunks=total_chunks,
                    user_id=user_id,
                )
                if not summary.strip():
                    return "error", None, "El modelo devolvió un resumen vacío para el fragmento."

                existing = (note.content or "").rstrip()
                separator = "\n\n" if existing else ""
                note.content = f"{existing}{separator}{summary.strip()}"
                note.ai_used = True
                if not commit_if_not_cancelled(job.id):
                    return "cancelled", None, None

                is_last = total_chunks and (chunk_index + 1) == total_chunks
                progress = f" ({chunk_index + 1}/{total_chunks})" if total_chunks > 1 else ""
                msg = f"Resumen en progreso{progress}: {note.title}"
                if is_last:
                    msg = f"Resumen completo listo: {note.title}"
                return "success", msg, None

            if job.type in ("flashcards_ai_new", "flashcards_ai_append"):
                payload = job.payload or {}
                user_id = payload.get("user_id")
                note_id = payload.get("note_id")
                deck_id = payload.get("deck_id")
                model = payload.get("model") or app.config["LMSTUDIO_MODEL"]
                custom_title = (payload.get("ai_deck_title") or "").strip()
                count = int(payload.get("count") or 5)

                note = Note.query.filter_by(id=note_id, user_id=user_id).first()
                if not note:
                    return "error", None, "Apunte/resumen inválido."

                cards = lmstudio_generate_flashcards(app, model, note, count=count)

                if job.type == "flashcards_ai_append":
                    deck = FlashcardDeck.query.filter_by(id=deck_id, user_id=user_id).first()
                    if not deck:
                        return "error", None, "Deck inválido."
                    deck.flashcards = (deck.flashcards or []) + cards
                    if not commit_if_not_cancelled(job.id):
                        return "cancelled", None, None
                    return "success", f"{len(cards)} flashcards añadidas a {deck.title}", None

                # new deck
                deck = FlashcardDeck(
                    user_id=user_id,
                    subject_id=note.subject_id,
                    title=custom_title or note.title,
                    exam_date=note.exam_date,
                    source_note_id=note.id,
                    flashcards=cards,
                )
                db.session.add(deck)
                if not commit_if_not_cancelled(job.id):
                    return "cancelled", None, None
                return "success", f"Deck creado: {deck.title}", None

            if job.type == "flashcards_ai_chunk":
                payload = job.payload or {}
                user_id = payload.get("user_id")
                note_id = payload.get("note_id")
                deck_id = payload.get("deck_id")
                model = payload.get("model") or app.config["LMSTUDIO_MODEL"]
                count = max(1, min(int(payload.get("count") or 5), 50))
                chunk_index = max(0, int(payload.get("chunk_index") or 0))
                total_chunks = max(1, int(payload.get("total_chunks") or 1))
                chunk_text = payload.get("text") or ""

                note = Note.query.filter_by(id=note_id, user_id=user_id).first()
                deck = FlashcardDeck.query.filter_by(id=deck_id, user_id=user_id).first()
                if not note or not deck:
                    return "error", None, "Apunte o deck inválido."

                cards = lmstudio_generate_flashcards(app, model, note, count=count, source_text=chunk_text)
                deck.flashcards = (deck.flashcards or []) + cards
                if not commit_if_not_cancelled(job.id):
                    return "cancelled", None, None

                msg = f"Flashcards añadidas (fragmento {chunk_index + 1}/{total_chunks}) a {deck.title}"
                if (chunk_index + 1) == total_chunks:
                    msg = f"Flashcards listas para {deck.title}"
                return "success", msg, None

            return "error", None, "Tipo de trabajo desconocido."
        except Timeout:
            db.session.rollback()
            return "error", None, "El modelo tardó demasiado en responder."
        except RequestException:
            db.session.rollback()
            return "error", None, "No se pudo conectar con LM Studio."
        except Exception as e:
            db.session.rollback()
            return "error", None, f"Error procesando el trabajo: {e}"


def create_app():
    app = Flask(__name__)

    os.makedirs(app.instance_path, exist_ok=True)
    upload_dir = os.getenv("UPLOAD_DIR") or os.path.join(app.instance_path, "uploads")
    os.makedirs(upload_dir, exist_ok=True)
    app.config["UPLOAD_DIR"] = upload_dir

    # LM Studio
    app.config["LMSTUDIO_API_BASE"] = os.getenv("LMSTUDIO_API_BASE", "http://127.0.0.1:1234/v1")
    app.config["LMSTUDIO_MODEL"] = os.getenv("LMSTUDIO_MODEL", "google/gemma-3-1b")
    app.config["LMSTUDIO_TIMEOUT"] = int(os.getenv("LMSTUDIO_TIMEOUT", "300"))
    app.config["ASK_PROFE_SESSION_SECONDS"] = int(os.getenv("ASK_PROFE_SESSION_SECONDS", "300"))

    # App/DB
    _secret = os.getenv("SECRET_KEY")
    if not _secret:
        _secret_file = os.getenv("SECRET_KEY_FILE", "/run/secrets/secret_key")
        if os.path.exists(_secret_file):
            with open(_secret_file) as _f:
                _secret = _f.read().strip()
    if not _secret:
        import secrets as _secrets
        import warnings
        _secret = _secrets.token_hex(32)
        warnings.warn(
            "SECRET_KEY not set — using ephemeral key. All sessions will reset on restart.",
            stacklevel=2,
        )
    app.config["SECRET_KEY"] = _secret
    db_uri = os.getenv("SQLALCHEMY_DATABASE_URI")
    if not db_uri:
        _db_host = os.getenv("DB_HOST")
        if _db_host:
            _pw = ""
            _pw_file = os.getenv("DB_PASSWORD_FILE")
            if _pw_file and os.path.exists(_pw_file):
                with open(_pw_file) as _f:
                    _pw = _f.read().strip()
            _db_user = os.getenv("DB_USER", "postgres")
            _db_port = os.getenv("DB_PORT", "5432")
            _db_name = os.getenv("DB_NAME", "postgres")
            db_uri = f"postgresql+psycopg2://{_db_user}:{_pw}@{_db_host}:{_db_port}/{_db_name}"
        else:
            db_uri = "sqlite:///app.db"
    app.config["SQLALCHEMY_DATABASE_URI"] = db_uri
    app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
    if db_uri.startswith("sqlite"):
        engine_options = app.config.setdefault("SQLALCHEMY_ENGINE_OPTIONS", {})
        connect_args = engine_options.setdefault("connect_args", {})
        connect_args.setdefault("timeout", 30)
        connect_args.setdefault("check_same_thread", False)
    else:
        engine_options = app.config.setdefault("SQLALCHEMY_ENGINE_OPTIONS", {})
        engine_options.setdefault("pool_pre_ping", True)

        @event.listens_for(Engine, "connect")
        def set_sqlite_pragma(dbapi_connection, _connection_record):
            if dbapi_connection.__class__.__module__.startswith("sqlite3"):
                cursor = dbapi_connection.cursor()
                cursor.execute("PRAGMA journal_mode=WAL;")
                cursor.execute("PRAGMA busy_timeout=30000;")
                cursor.close()

    db.init_app(app)
    migrate.init_app(app, db)

    CSRFProtect(app)
    limiter.init_app(app)

    login_manager = LoginManager()
    login_manager.login_view = "login"
    login_manager.init_app(app)

    @app.errorhandler(429)
    def ratelimit_handler(e):
        flash("Demasiados intentos. Espera un momento antes de volver a intentarlo.", "danger")
        return redirect(request.referrer or url_for("login")), 429

    @app.errorhandler(RequestEntityTooLarge)
    def handle_file_too_large(_error):
        max_kb = (app.config.get("MAX_CONTENT_LENGTH") or MAX_UPLOAD_BYTES) // 1024
        flash(f"Archivo demasiado grande. Máximo {max_kb} KB.", "error")
        return redirect(request.referrer or url_for("dashboard"))

    def queue_has_work() -> bool:
        return db.session.query(Job.id).filter(Job.status.in_(("pending", "running"))).first() is not None

    def get_active_profe_lock(now: datetime | None = None) -> ProfeSessionLock | None:
        now = now or datetime.utcnow()
        lock = ProfeSessionLock.query.order_by(ProfeSessionLock.ends_at.desc()).first()
        if not lock or not lock.ends_at:
            return None
        if lock.ends_at <= now:
            return None
        return lock

    def start_profe_lock(user_id: int, now: datetime | None = None) -> ProfeSessionLock:
        now = now or datetime.utcnow()
        ends_at = now + timedelta(seconds=app.config["ASK_PROFE_SESSION_SECONDS"])
        lock = ProfeSessionLock.query.order_by(ProfeSessionLock.ends_at.desc()).first()
        if lock:
            lock.user_id = user_id
            lock.starts_at = now
            lock.ends_at = ends_at
        else:
            lock = ProfeSessionLock(user_id=user_id, starts_at=now, ends_at=ends_at)
            db.session.add(lock)
        db.session.commit()
        return lock

    def profe_is_busy(user_id: int | None = None) -> bool:
        lock = get_active_profe_lock()
        if lock:
            if user_id is not None and lock.user_id == user_id:
                return False
            return True
        return queue_has_work()

    @app.context_processor
    def inject_login_flag():
        # Flags to trigger mascot celebration on first render after login/register.
        just_logged_in = session.pop("just_logged_in", None)
        just_registered = session.pop("just_registered", None)
        return {"just_logged_in": bool(just_logged_in), "just_registered": bool(just_registered)}

    @app.context_processor
    def inject_profe_busy_flag():
        if not current_user.is_authenticated:
            return {"profe_busy": False, "llm_busy": False}
        return {"profe_busy": profe_is_busy(current_user.id), "llm_busy": queue_has_work()}

    @app.context_processor
    def inject_deck_helpers():
        def deck_display_title(title: str | None, exam_date) -> str:
            if not title:
                return ""
            if exam_date:
                suffix = f" ({exam_date})"
                if title.endswith(suffix):
                    return title[: -len(suffix)].rstrip()
            return title

        return {"deck_display_title": deck_display_title}

    @app.before_request
    def enforce_setup_completion():
        endpoint = request.endpoint or ""
        if endpoint in ("login", "register", "logout", "setup", "static"):
            return None
        if not current_user.is_authenticated:
            return None
        if not is_setup_complete(current_user.id):
            return redirect(url_for("setup"))
        return None

    @login_manager.user_loader
    def load_user(user_id):
        return User.query.get(int(user_id))

    def acquire_process_lock(name: str, blocking: bool = True):
        lock_path = os.path.join(tempfile.gettempdir(), name)
        lock_file = open(lock_path, "a+")
        flags = fcntl.LOCK_EX
        if not blocking:
            flags |= fcntl.LOCK_NB
        try:
            fcntl.flock(lock_file.fileno(), flags)
        except BlockingIOError:
            lock_file.close()
            return None
        return lock_file

    with app.app_context():
        global _worker_started, _worker_lock_handle
        if not _worker_started:
            if _worker_lock_handle is None:
                _worker_lock_handle = acquire_process_lock("flask-devops-demo-worker.lock", blocking=False)
            if _worker_lock_handle:
                def worker_loop():
                    retry_job_id = None
                    retry_at = None
                    while True:
                        with app.app_context():
                            try:
                                if get_active_profe_lock():
                                    time.sleep(2)
                                    continue

                                job = None
                                if retry_job_id:
                                    now = time.time()
                                    if retry_at and now < retry_at:
                                        time.sleep(min(2, retry_at - now))
                                        continue
                                    job = Job.query.filter_by(id=retry_job_id).first()
                                    if not job:
                                        retry_job_id = None
                                        retry_at = None
                                        time.sleep(1)
                                        continue
                                if not job:
                                    job = (
                                        Job.query.filter_by(status="pending")
                                        .order_by(Job.created_at.asc(), Job.id.asc())
                                        .first()
                                    )
                                    if not job:
                                        time.sleep(2)
                                        continue

                                job.status = "running"
                                db.session.commit()

                                status, msg, err = process_job(app, job)
                                db.session.refresh(job)
                                if job.status == "cancelled":
                                    retry_job_id = None
                                    retry_at = None
                                    continue
                                if status == "error":
                                    job.status = "pending"
                                    job.result_message = None
                                    job.error_message = err
                                    job.updated_at = datetime.utcnow()
                                    db.session.commit()
                                    retry_job_id = job.id
                                    retry_at = time.time() + JOB_RETRY_SECONDS
                                    continue

                                job.status = status
                                job.result_message = msg
                                job.error_message = err
                                job.updated_at = datetime.utcnow()
                                db.session.commit()
                                retry_job_id = None
                                retry_at = None
                            except OperationalError as exc:
                                db.session.rollback()
                                if "database is locked" in str(exc).lower():
                                    time.sleep(1.5)
                                    continue
                                app.logger.exception("Database error in worker loop")
                                time.sleep(2)
                                continue
                        # pequeña pausa para no saturar
                        time.sleep(0.5)

                threading.Thread(target=worker_loop, daemon=True).start()
                _worker_started = True

    # ---------- AUTH ----------
    @app.route("/", methods=["GET", "POST"])
    @limiter.limit("10 per minute", methods=["POST"])
    def login():
        if current_user.is_authenticated:
            if not is_setup_complete(current_user.id):
                return redirect(url_for("setup"))
            return redirect(url_for("dashboard"))

        if request.method == "POST":
            username = (request.form.get("username") or "").strip()
            password = request.form.get("password") or ""
            user = User.query.filter_by(username=username).first()

            if user and password and check_password_hash(user.password_hash, password):
                login_user(user)
                flash("Sesión iniciada ✅", "login_success")
                session["just_logged_in"] = True
                if not is_setup_complete(user.id):
                    return redirect(url_for("setup"))
                return redirect(url_for("dashboard"))

            flash("Usuario o contraseña incorrectos.", "error")
        return render_template("login.html")

    @app.route("/register", methods=["GET", "POST"])
    @limiter.limit("5 per hour", methods=["POST"])
    def register():
        if current_user.is_authenticated:
            if not is_setup_complete(current_user.id):
                return redirect(url_for("setup"))
            return redirect(url_for("dashboard"))

        if request.method == "POST":
            username = (request.form.get("username") or "").strip()
            password = request.form.get("password") or ""

            if not username or not password:
                flash("Rellena usuario y contraseña.", "error")
                return redirect(url_for("register"))

            if User.query.filter_by(username=username).first():
                flash("Ese usuario ya existe.", "error")
                return redirect(url_for("register"))

            hashed_pw = generate_password_hash(password, method="pbkdf2:sha256", salt_length=16)
            new_user = User(username=username, password_hash=hashed_pw)
            db.session.add(new_user)
            db.session.commit()

            login_user(new_user)
            flash("Cuenta creada e inicio de sesión ✅", "login_success")
            session["just_logged_in"] = True
            session["just_registered"] = True
            return redirect(url_for("setup"))

        return render_template("register.html")

    @app.route("/setup", methods=["GET", "POST"])
    @login_required
    def setup():
        profile = StudentProfile.query.filter_by(user_id=current_user.id).first()
        student_name = profile.student_name if profile else ""
        student_age = ""
        personality_notes = profile.personality_notes if profile else ""
        llm_profile = (
            profile.llm_profile
            if profile and profile.llm_profile in LLM_PROFILE_PRESETS
            else DEFAULT_LLM_PROFILE
        )

        if request.method == "POST":
            student_name = (request.form.get("student_name") or "").strip()
            student_age = (request.form.get("student_age") or "").strip()
            personality_notes = (request.form.get("personality_notes") or "").strip()
            llm_profile = (request.form.get("llm_profile") or llm_profile).strip()

            errors: list[str] = []
            if not student_name:
                errors.append("Escribe el nombre del estudiante.")

            age_val = None
            birth_date = None
            if student_age:
                if "-" in student_age:
                    try:
                        birth_date = datetime.strptime(student_age, "%Y-%m-%d").date()
                    except ValueError:
                        birth_date = None
                    if birth_date:
                        today = datetime.utcnow().date()
                        age_val = today.year - birth_date.year
                        if (today.month, today.day) < (birth_date.month, birth_date.day):
                            age_val -= 1
                else:
                    try:
                        age_val = int(student_age)
                    except (TypeError, ValueError):
                        age_val = None

            if not age_val or age_val < 1:
                errors.append("Indica una fecha de nacimiento válida.")

            if not personality_notes:
                errors.append("Añade detalles de personalidad para contextualizar al profe.")
            if llm_profile not in LLM_PROFILE_PRESETS:
                errors.append("Selecciona un perfil de VRAM válido.")

            if errors:
                for err in errors:
                    flash(err, "error")
                return render_template(
                    "setup.html",
                    student_name=student_name,
                    student_age=student_age,
                    personality_notes=personality_notes,
                    llm_profiles=LLM_PROFILE_CHOICES,
                    llm_profile=llm_profile,
                    age_label="Fecha de nacimiento",
                    age_mode="dob",
                )

            try:
                profile = StudentProfile.query.filter_by(user_id=current_user.id).first()
                if profile:
                    profile.student_name = student_name
                    profile.age = age_val or 0
                    profile.personality_notes = personality_notes
                    profile.llm_profile = llm_profile
                else:
                    profile = StudentProfile(
                        user_id=current_user.id,
                        student_name=student_name,
                        age=age_val or 0,
                        personality_notes=personality_notes,
                        llm_profile=llm_profile,
                        task_window_days=TASK_WINDOW_DEFAULT,
                    )
                    db.session.add(profile)

                db.session.commit()
                flash("Datos guardados ✅", "success")
                return redirect(url_for("setup_subjects"))
            except Exception:
                db.session.rollback()
                flash("Error guardando la configuración.", "error")
                return render_template(
                    "setup.html",
                    student_name=student_name,
                    student_age=student_age,
                    personality_notes=personality_notes,
                    llm_profiles=LLM_PROFILE_CHOICES,
                    llm_profile=llm_profile,
                    age_label="Fecha de nacimiento",
                    age_mode="dob",
                )

        return render_template(
            "setup.html",
            student_name=student_name,
            student_age=student_age,
            personality_notes=personality_notes,
            llm_profiles=LLM_PROFILE_CHOICES,
            llm_profile=llm_profile,
            age_label="Fecha de nacimiento",
            age_mode="dob",
        )

    @app.route("/setup/subjects", methods=["GET", "POST"])
    @login_required
    def setup_subjects():
        if not is_setup_complete(current_user.id):
            return redirect(url_for("setup"))

        subjects_seed: list[dict] = []

        if request.method == "POST":
            subjects_payload = request.form.get("subjects_payload") or ""
            subjects_clean: list[dict] = []
            errors: list[str] = []

            raw_subjects = []
            if subjects_payload:
                try:
                    raw_subjects = json.loads(subjects_payload)
                except (TypeError, ValueError):
                    raw_subjects = None
                    errors.append("No se pudo leer la lista de asignaturas.")

            if raw_subjects is not None:
                if not isinstance(raw_subjects, list):
                    errors.append("Formato inválido de asignaturas.")
                else:
                    seen_names: set[str] = set()
                    for subj in raw_subjects:
                        if not isinstance(subj, dict):
                            continue
                        raw_subject_id = subj.get("id")
                        subject_id = None
                        if raw_subject_id is not None:
                            try:
                                subject_id = int(raw_subject_id)
                            except (TypeError, ValueError):
                                subject_id = None
                        name = (subj.get("name") or "").strip()
                        color_raw = (subj.get("color") or "").strip()
                        color = normalize_subject_color(color_raw)
                        exams = subj.get("exams") if isinstance(subj, dict) else []
                        if not isinstance(exams, list):
                            exams = []
                        subject_seed = {
                            "id": subject_id,
                            "name": name,
                            "color": color or color_raw,
                            "exams": [],
                        }

                        subject_has_content = bool(name or color_raw)
                        for exam in exams:
                            raw_exam_id = exam.get("id") if isinstance(exam, dict) else None
                            exam_id = None
                            if raw_exam_id is not None:
                                try:
                                    exam_id = int(raw_exam_id)
                                except (TypeError, ValueError):
                                    exam_id = None
                            date_str = (exam.get("date") or "").strip() if isinstance(exam, dict) else ""
                            tema = (exam.get("tema") or "").strip() if isinstance(exam, dict) else ""
                            if date_str or tema:
                                subject_has_content = True
                            subject_seed["exams"].append({"id": exam_id, "date": date_str, "tema": tema})

                        if not subject_has_content:
                            continue

                        subjects_seed.append(subject_seed)

                        if not name:
                            errors.append("Cada asignatura necesita un nombre.")
                            continue
                        if color_raw and not color:
                            errors.append(f'Color inválido en "{name}".')
                            continue
                        lowered = name.lower()
                        if lowered in seen_names:
                            errors.append(f"La asignatura \"{name}\" está duplicada.")
                            continue
                        seen_names.add(lowered)

                        exam_entries: list[dict] = []
                        seen_exam_keys: set[tuple] = set()
                        for exam in exams:
                            raw_exam_id = exam.get("id") if isinstance(exam, dict) else None
                            exam_id = None
                            if raw_exam_id is not None:
                                try:
                                    exam_id = int(raw_exam_id)
                                except (TypeError, ValueError):
                                    exam_id = None
                            date_str = (exam.get("date") or "").strip() if isinstance(exam, dict) else ""
                            tema = (exam.get("tema") or "").strip() if isinstance(exam, dict) else ""
                            if not date_str and not tema:
                                continue
                            if not date_str or not tema:
                                errors.append(
                                    f"Cada examen de \"{name}\" debe tener fecha y tema."
                                )
                                continue
                            try:
                                exam_date = datetime.strptime(date_str, "%Y-%m-%d").date()
                            except ValueError:
                                errors.append(
                                    f"Formato de fecha inválido en \"{name}\": {date_str}."
                                )
                                continue
                            exam_key = (exam_date, tema.lower())
                            if exam_key in seen_exam_keys:
                                errors.append(
                                    f"El examen \"{tema}\" de \"{name}\" está duplicado."
                                )
                                continue
                            seen_exam_keys.add(exam_key)
                            exam_entries.append({"id": exam_id, "date": exam_date, "tema": tema})

                        subjects_clean.append(
                            {"id": subject_id, "name": name, "color": color, "exams": exam_entries}
                        )

            if errors:
                for err in errors:
                    flash(err, "error")
                return render_template("setup_subjects.html", subjects_seed=subjects_seed)

            try:
                existing_subjects = Subject.query.filter_by(user_id=current_user.id).all()
                existing_subjects_by_id = {s.id: s for s in existing_subjects}
                incoming_subject_ids: set[int] = set()

                def delete_exam_content(subject_id: int, exam_date):
                    note_ids = [
                        row[0]
                        for row in db.session.query(Note.id)
                        .filter_by(user_id=current_user.id, subject_id=subject_id, exam_date=exam_date)
                        .all()
                    ]
                    if note_ids:
                        NoteSourceFile.query.filter(NoteSourceFile.note_id.in_(note_ids)).delete(
                            synchronize_session=False
                        )
                    Note.query.filter_by(
                        user_id=current_user.id, subject_id=subject_id, exam_date=exam_date
                    ).delete(synchronize_session=False)
                    FlashcardDeck.query.filter_by(
                        user_id=current_user.id, subject_id=subject_id, exam_date=exam_date
                    ).delete(synchronize_session=False)

                for subj in subjects_clean:
                    subject_id = subj.get("id")
                    subject = existing_subjects_by_id.get(subject_id) if subject_id else None
                    if subject:
                        subject.name = subj["name"]
                        subject.color = subj.get("color")
                    else:
                        subject = Subject(
                            user_id=current_user.id,
                            name=subj["name"],
                            color=subj.get("color"),
                        )
                        db.session.add(subject)
                        db.session.flush()
                    incoming_subject_ids.add(subject.id)

                    existing_exams = {e.id: e for e in SubjectExam.query.filter_by(subject_id=subject.id).all()}
                    incoming_exam_ids: set[int] = set()

                    for exam in subj["exams"]:
                        exam_id = exam.get("id")
                        if exam_id and exam_id in existing_exams:
                            exam_row = existing_exams[exam_id]
                            old_date = exam_row.exam_date
                            exam_row.exam_date = exam["date"]
                            exam_row.tema = exam["tema"]
                            incoming_exam_ids.add(exam_row.id)
                            if old_date != exam_row.exam_date:
                                Note.query.filter_by(
                                    user_id=current_user.id,
                                    subject_id=subject.id,
                                    exam_date=old_date,
                                ).update({Note.exam_date: exam_row.exam_date}, synchronize_session=False)
                                FlashcardDeck.query.filter_by(
                                    user_id=current_user.id,
                                    subject_id=subject.id,
                                    exam_date=old_date,
                                ).update({FlashcardDeck.exam_date: exam_row.exam_date}, synchronize_session=False)
                        else:
                            new_exam = SubjectExam(
                                subject_id=subject.id,
                                exam_date=exam["date"],
                                tema=exam["tema"],
                            )
                            db.session.add(new_exam)
                            db.session.flush()
                            incoming_exam_ids.add(new_exam.id)

                    for exam_id, exam_row in existing_exams.items():
                        if exam_id not in incoming_exam_ids:
                            delete_exam_content(subject.id, exam_row.exam_date)
                            db.session.delete(exam_row)

                for subject in existing_subjects:
                    if subject.id not in incoming_subject_ids:
                        db.session.delete(subject)

                db.session.commit()
                flash("Asignaturas guardadas ✅", "success")
                return redirect(url_for("setup_generate"))
            except Exception:
                db.session.rollback()
                flash("Error guardando las asignaturas.", "error")
                return render_template("setup_subjects.html", subjects_seed=subjects_seed)

        if request.method == "GET":
            subjects = (
                Subject.query.filter_by(user_id=current_user.id)
                .order_by(Subject.name.asc())
                .all()
            )
            subjects_seed = []
            for subject in subjects:
                exams = (
                    SubjectExam.query.filter_by(subject_id=subject.id)
                    .order_by(SubjectExam.exam_date.asc(), SubjectExam.tema.asc())
                    .all()
                )
                subjects_seed.append(
                    {
                        "id": subject.id,
                        "name": subject.name,
                        "color": subject.color,
                        "exams": [
                            {
                                "id": exam.id,
                                "date": exam.exam_date.isoformat() if exam.exam_date else "",
                                "tema": exam.tema,
                            }
                            for exam in exams
                        ],
                    }
                )
        return render_template("setup_subjects.html", subjects_seed=subjects_seed)

    @app.route("/setup/generate", methods=["GET", "POST"])
    @login_required
    def setup_generate():
        if not is_setup_complete(current_user.id):
            return redirect(url_for("setup"))

        profile = StudentProfile.query.filter_by(user_id=current_user.id).first()
        available_models = fetch_models(app)
        selected_model = resolve_default_model(app, current_user.id, available_models)

        subjects = Subject.query.filter_by(user_id=current_user.id).order_by(Subject.name.asc()).all()
        subject_blocks: list[dict] = []
        for subject in subjects:
            exams = (
                SubjectExam.query.filter_by(subject_id=subject.id)
                .order_by(SubjectExam.exam_date.asc(), SubjectExam.tema.asc())
                .all()
            )
            if exams:
                subject_blocks.append({"subject": subject, "exams": exams})

        if request.method == "POST":
            selected_model = request.form.get("default_model") or selected_model
            if available_models and selected_model not in available_models:
                selected_model = available_models[0]
            if profile:
                profile.default_model = selected_model
                db.session.commit()

            total_files = 0
            queued_files = 0
            errors: list[str] = []
            upload_dir = app.config.get("UPLOAD_DIR") or app.instance_path

            try:
                for block in subject_blocks:
                    subject = block["subject"]
                    for exam in block["exams"]:
                        uploads = request.files.getlist(f"files_{exam.id}")
                        for upload in uploads:
                            if not upload or not upload.filename:
                                continue
                            total_files += 1
                            if not allowed_file(upload.filename):
                                errors.append(f"{upload.filename}: solo se permiten archivos .txt, .pdf o .pptx.")
                                continue

                            filename = secure_filename(upload.filename)
                            if not filename:
                                errors.append("Nombre de archivo inválido.")
                                continue

                            stored_name = f"{uuid.uuid4().hex}_{filename}"
                            stored_path = os.path.join(upload_dir, stored_name)
                            try:
                                size, err = save_upload_stream(upload, stored_path, MAX_UPLOAD_BYTES)
                            except Exception:
                                errors.append(f"{filename}: error guardando el archivo.")
                                if os.path.exists(stored_path):
                                    os.remove(stored_path)
                                continue
                            if err == "empty":
                                errors.append(f"{filename}: el archivo está vacío.")
                                continue
                            if err == "too_large":
                                errors.append(f"{filename}: archivo demasiado grande (máx {MAX_UPLOAD_BYTES // 1024} KB).")
                                continue
                            if size <= 0:
                                errors.append(f"{filename}: el archivo está vacío.")
                                continue

                            db.session.add(
                                Job(
                                    user_id=current_user.id,
                                    type="file_import",
                                    payload={
                                        "user_id": current_user.id,
                                        "subject_id": subject.id,
                                        "exam_id": exam.id,
                                        "filename": filename,
                                        "file_path": stored_path,
                                        "content_type": upload.mimetype or "",
                                        "model": selected_model,
                                    },
                                )
                            )
                            queued_files += 1

                if queued_files:
                    db.session.commit()
                    flash(
                        f"Archivos en cola ✅ {queued_files} archivo(s). Se procesarán en segundo plano.",
                        "success",
                    )
                else:
                    db.session.rollback()
                    if total_files == 0:
                        flash("No has subido archivos. Puedes hacerlo más tarde.", "success")
            except Exception:
                db.session.rollback()
                flash("Error procesando la generación inicial.", "error")
                return redirect(url_for("setup_generate"))

            for err in errors:
                flash(err, "error")
            return redirect(url_for("dashboard"))

        return render_template(
            "setup_generate.html",
            subjects=subject_blocks,
            models=available_models,
            selected_model=selected_model,
            max_kb=MAX_UPLOAD_BYTES // 1024,
        )

    @app.route("/setup/next")
    @login_required
    def setup_next():
        if not is_setup_complete(current_user.id):
            return redirect(url_for("setup"))
        return redirect(url_for("dashboard"))

    @app.route("/logout")
    @login_required
    def logout():
        logout_user()
        return redirect(url_for("login"))

    # ---------- DASHBOARD ----------
    @app.route("/dashboard")
    @login_required
    def dashboard():
        profile = StudentProfile.query.filter_by(user_id=current_user.id).first()
        student_name = profile.student_name if profile else current_user.username
        window_days, window_start, window_end = calendar_window_range(profile)
        calendar_items = build_calendar_items(current_user.id, window_start, window_end, window_start)
        calendar_label = TASK_WINDOW_LABELS.get(window_days, f"{window_days} días")
        queue_groups = build_queue_groups(current_user.id)
        return render_template(
            "dashboard.html",
            student_name=student_name,
            queue_groups=queue_groups,
            calendar_items=calendar_items,
            calendar_window_label=calendar_label,
            calendar_window_start=window_start,
            calendar_window_end=window_end,
        )

    @app.route("/calendar")
    @login_required
    def task_calendar():
        profile = StudentProfile.query.filter_by(user_id=current_user.id).first()
        student_name = profile.student_name if profile else current_user.username
        window_days, window_start, window_end = calendar_window_range(profile)
        calendar_items = build_calendar_items(current_user.id, window_start, window_end, window_start)
        window_label = TASK_WINDOW_LABELS.get(window_days, f"{window_days} días")
        subjects = Subject.query.filter_by(user_id=current_user.id).order_by(Subject.name.asc()).all()
        return render_template(
            "calendar.html",
            student_name=student_name,
            window_options=TASK_WINDOW_OPTIONS,
            window_days=window_days,
            window_label=window_label,
            window_start=window_start,
            window_end=window_end,
            calendar_items=calendar_items,
            subjects=subjects,
        )

    @app.route("/calendar/window", methods=["POST"])
    @login_required
    def task_calendar_window_update():
        window_days = request.form.get("window_days", type=int)
        if window_days not in TASK_WINDOW_VALUES:
            flash("Selecciona un periodo válido.", "error")
            return redirect(url_for("task_calendar"))

        profile = StudentProfile.query.filter_by(user_id=current_user.id).first()
        if not profile:
            flash("Completa el perfil antes de configurar el calendario.", "error")
            return redirect(url_for("setup"))

        try:
            profile.task_window_days = window_days
            db.session.commit()
            flash("Periodo actualizado ✅", "success")
        except Exception:
            db.session.rollback()
            flash("No se pudo actualizar el periodo.", "error")
        return redirect(url_for("task_calendar"))

    @app.route("/calendar/tasks", methods=["POST"])
    @login_required
    def calendar_task_create():
        title = (request.form.get("title") or "").strip()
        due_date_str = (request.form.get("due_date") or "").strip()
        notes = (request.form.get("notes") or "").strip()
        subject_id = request.form.get("subject_id", type=int)
        errors: list[str] = []

        if not title:
            errors.append("Añade un título para la tarea.")

        due_date = None
        if not due_date_str:
            errors.append("Selecciona una fecha límite.")
        else:
            try:
                due_date = datetime.strptime(due_date_str, "%Y-%m-%d").date()
            except ValueError:
                errors.append("Formato de fecha inválido.")

        subject = None
        if subject_id:
            subject = Subject.query.filter_by(id=subject_id, user_id=current_user.id).first()
            if not subject:
                errors.append("Asignatura inválida.")

        if errors:
            for err in errors:
                flash(err, "error")
            return redirect(url_for("task_calendar"))

        try:
            task = TaskItem(
                user_id=current_user.id,
                subject_id=subject.id if subject else None,
                title=title,
                due_date=due_date,
                notes=notes or None,
            )
            db.session.add(task)
            db.session.commit()
            flash("Tarea añadida ✅", "success")
        except Exception:
            db.session.rollback()
            flash("No se pudo guardar la tarea.", "error")
        return redirect(url_for("task_calendar"))

    @app.route("/calendar/exams", methods=["POST"])
    @login_required
    def calendar_exam_create():
        subject_id = request.form.get("subject_id", type=int)
        exam_date_str = (request.form.get("exam_date") or "").strip()
        tema = (request.form.get("tema") or "").strip()
        new_subject_name = (request.form.get("new_subject") or "").strip()
        new_subject_color_raw = (request.form.get("new_subject_color") or "").strip()
        errors: list[str] = []

        subject = None
        if new_subject_name:
            existing_subject = Subject.query.filter_by(
                user_id=current_user.id,
                name=new_subject_name,
            ).first()
            if existing_subject:
                subject = existing_subject
            else:
                if len(new_subject_name) > 120:
                    errors.append("El nombre de la asignatura es demasiado largo.")
                color_value = None
                if new_subject_color_raw:
                    color_value = normalize_subject_color(new_subject_color_raw)
                    if not color_value:
                        errors.append("Color de asignatura inválido.")
                if not errors:
                    subject = Subject(
                        user_id=current_user.id,
                        name=new_subject_name,
                        color=color_value,
                    )
                    db.session.add(subject)
                    db.session.flush()
        else:
            if not subject_id:
                errors.append("Selecciona una asignatura.")
            else:
                subject = Subject.query.filter_by(id=subject_id, user_id=current_user.id).first()
                if not subject:
                    errors.append("Asignatura inválida.")

        exam_date = None
        if not exam_date_str:
            errors.append("Selecciona una fecha de examen.")
        else:
            try:
                exam_date = datetime.strptime(exam_date_str, "%Y-%m-%d").date()
            except ValueError:
                errors.append("Formato de fecha inválido.")

        if not tema:
            errors.append("Indica el tema del examen.")

        if errors:
            db.session.rollback()
            for err in errors:
                flash(err, "error")
            return redirect(url_for("task_calendar"))

        try:
            existing = SubjectExam.query.filter_by(
                subject_id=subject.id,
                exam_date=exam_date,
                tema=tema,
            ).first()
            if existing:
                flash("Ese examen ya existe.", "error")
                return redirect(url_for("task_calendar"))

            new_exam = SubjectExam(subject_id=subject.id, exam_date=exam_date, tema=tema)
            db.session.add(new_exam)
            db.session.commit()
            flash("Examen añadido ✅", "success")
        except Exception:
            db.session.rollback()
            flash("No se pudo añadir el examen.", "error")
        return redirect(url_for("task_calendar"))

    @app.route("/calendar/tasks/<int:task_id>/update", methods=["POST"])
    @login_required
    def calendar_task_update(task_id: int):
        task = TaskItem.query.filter_by(id=task_id, user_id=current_user.id).first()
        if not task:
            flash("Tarea no encontrada.", "error")
            return redirect(url_for("task_calendar"))

        title = (request.form.get("title") or "").strip()
        due_date_str = (request.form.get("due_date") or "").strip()
        notes = (request.form.get("notes") or "").strip()
        subject_id = request.form.get("subject_id", type=int)
        errors: list[str] = []

        if not title:
            errors.append("El título no puede estar vacío.")

        due_date = None
        if not due_date_str:
            errors.append("Selecciona una fecha límite.")
        else:
            try:
                due_date = datetime.strptime(due_date_str, "%Y-%m-%d").date()
            except ValueError:
                errors.append("Formato de fecha inválido.")

        subject = None
        if subject_id:
            subject = Subject.query.filter_by(id=subject_id, user_id=current_user.id).first()
            if not subject:
                errors.append("Asignatura inválida.")

        if errors:
            for err in errors:
                flash(err, "error")
            return redirect(url_for("task_calendar"))

        try:
            task.title = title
            task.due_date = due_date
            task.notes = notes or None
            task.subject_id = subject.id if subject else None
            db.session.commit()
            flash("Tarea actualizada ✅", "success")
        except Exception:
            db.session.rollback()
            flash("No se pudo actualizar la tarea.", "error")
        return redirect(url_for("task_calendar"))

    @app.route("/calendar/tasks/<int:task_id>/delete", methods=["POST"])
    @login_required
    def calendar_task_delete(task_id: int):
        task = TaskItem.query.filter_by(id=task_id, user_id=current_user.id).first()
        if not task:
            flash("Tarea no encontrada.", "error")
            return redirect(url_for("task_calendar"))
        try:
            db.session.delete(task)
            db.session.commit()
            flash("Tarea eliminada ✅", "success")
        except Exception:
            db.session.rollback()
            flash("No se pudo eliminar la tarea.", "error")
        return redirect(url_for("task_calendar"))

    @app.route("/jobs/cancel", methods=["POST"])
    @login_required
    def cancel_jobs():
        group_kind = (request.form.get("group_kind") or "").strip()
        if group_kind == "note":
            note_id = request.form.get("note_id", type=int)
            if not note_id:
                flash("Apunte inválido para cancelar.", "error")
                return redirect(url_for("dashboard"))

            jobs = (
                Job.query.filter(
                    Job.user_id == current_user.id,
                    Job.type.in_(
                        (
                            "note_ai",
                            "note_ai_chunk",
                            "flashcards_ai_new",
                            "flashcards_ai_append",
                            "flashcards_ai_chunk",
                        )
                    ),
                )
                .order_by(Job.created_at.asc(), Job.id.asc())
                .all()
            )
            related_jobs = [job for job in jobs if payload_int(job.payload or {}, "note_id") == note_id]
            if not related_jobs:
                flash("No se encontraron trabajos para ese apunte.", "error")
                return redirect(url_for("dashboard"))

            incomplete_jobs = [job for job in related_jobs if job.status in INCOMPLETE_JOB_STATUSES]
            if not incomplete_jobs:
                flash("No hay trabajos en curso para ese apunte.", "error")
                return redirect(url_for("dashboard"))
            start_at = min((job.created_at or datetime.min) for job in incomplete_jobs)
            related_jobs = [
                job
                for job in related_jobs
                if (job.created_at or datetime.min) >= start_at
            ]
            deck_ids = {
                payload_int(job.payload or {}, "deck_id")
                for job in related_jobs
                if payload_int(job.payload or {}, "deck_id") is not None
            }

            for job in related_jobs:
                if job.status in INCOMPLETE_JOB_STATUSES:
                    job.status = "cancelled"
                    job.result_message = None
                    job.error_message = "Cancelado por el usuario."
                    job.updated_at = datetime.utcnow()

            has_note_generation = any(
                job.type in ("note_ai", "note_ai_chunk") for job in related_jobs
            )

            decks_to_check = []
            if deck_ids:
                decks_to_check = (
                    FlashcardDeck.query.filter(
                        FlashcardDeck.user_id == current_user.id,
                        FlashcardDeck.id.in_(deck_ids),
                    )
                    .all()
                )

            for deck in decks_to_check:
                if has_note_generation and deck.source_note_id == note_id:
                    db.session.delete(deck)
                    continue
                deck_jobs = [
                    job
                    for job in related_jobs
                    if payload_int(job.payload or {}, "deck_id") == deck.id
                ]
                sizes = [
                    payload_int(job.payload or {}, "deck_size_before")
                    for job in deck_jobs
                    if payload_int(job.payload or {}, "deck_size_before") is not None
                ]
                if sizes:
                    keep_size = min(sizes)
                    deck.flashcards = (deck.flashcards or [])[:keep_size]

            if has_note_generation:
                note = Note.query.filter_by(id=note_id, user_id=current_user.id).first()
                if note:
                    NoteSourceFile.query.filter_by(note_id=note.id).delete()
                    db.session.delete(note)

            db.session.commit()
            flash("Generación cancelada ✅ Se han eliminado resultados parciales.", "success")
            return redirect(url_for("dashboard"))

        if group_kind == "file_import":
            job_id = request.form.get("job_id", type=int)
            if not job_id:
                flash("Trabajo inválido para cancelar.", "error")
                return redirect(url_for("dashboard"))
            job = Job.query.filter_by(id=job_id, user_id=current_user.id, type="file_import").first()
            if not job:
                flash("Trabajo no encontrado.", "error")
                return redirect(url_for("dashboard"))
            if job.status in INCOMPLETE_JOB_STATUSES:
                job.status = "cancelled"
                job.result_message = None
                job.error_message = "Cancelado por el usuario."
                job.updated_at = datetime.utcnow()
            payload = job.payload or {}
            file_path = payload.get("file_path") if isinstance(payload, dict) else None
            if file_path and os.path.exists(file_path):
                try:
                    os.remove(file_path)
                except OSError:
                    pass
            db.session.commit()
            flash("Trabajo cancelado ✅", "success")
            return redirect(url_for("dashboard"))

        flash("No se pudo cancelar la cola solicitada.", "error")
        return redirect(url_for("dashboard"))

    @app.route("/options")
    @login_required
    def options():
        return render_template("options.html")

    @app.route("/options/profile", methods=["GET", "POST"])
    @login_required
    def profile_edit():
        profile = StudentProfile.query.filter_by(user_id=current_user.id).first()
        if not profile:
            return redirect(url_for("setup"))

        student_name = profile.student_name or ""
        student_age = str(profile.age or "")
        personality_notes = profile.personality_notes or ""
        llm_profile = profile.llm_profile if profile.llm_profile in LLM_PROFILE_PRESETS else DEFAULT_LLM_PROFILE

        if request.method == "POST":
            student_name = (request.form.get("student_name") or "").strip()
            student_age = (request.form.get("student_age") or "").strip()
            personality_notes = (request.form.get("personality_notes") or "").strip()
            llm_profile = (request.form.get("llm_profile") or llm_profile).strip()

            errors: list[str] = []
            if not student_name:
                errors.append("Escribe el nombre del estudiante.")

            age_val = None
            try:
                age_val = int(student_age)
                if age_val < 1:
                    raise ValueError
            except (TypeError, ValueError):
                errors.append("Indica una edad válida.")

            if not personality_notes:
                errors.append("Añade detalles de personalidad para contextualizar al profe.")
            if llm_profile not in LLM_PROFILE_PRESETS:
                errors.append("Selecciona un perfil de VRAM válido.")

            if errors:
                for err in errors:
                    flash(err, "error")
                return render_template(
                    "setup.html",
                    student_name=student_name,
                    student_age=student_age,
                    personality_notes=personality_notes,
                    llm_profiles=LLM_PROFILE_CHOICES,
                    llm_profile=llm_profile,
                    page_title="Perfil del estudiante",
                    heading="Perfil del estudiante",
                    subtext="Actualiza los datos del estudiante cuando lo necesites.",
                    submit_label="Guardar cambios",
                    back_url=url_for("options"),
                )

            profile.student_name = student_name
            profile.age = age_val or 0
            profile.personality_notes = personality_notes
            profile.llm_profile = llm_profile
            db.session.commit()
            flash("Perfil actualizado ✅", "success")
            return redirect(url_for("options"))

        return render_template(
            "setup.html",
            student_name=student_name,
            student_age=student_age,
            personality_notes=personality_notes,
            llm_profiles=LLM_PROFILE_CHOICES,
            llm_profile=llm_profile,
            page_title="Perfil del estudiante",
            heading="Perfil del estudiante",
            subtext="Actualiza los datos del estudiante cuando lo necesites.",
            submit_label="Guardar cambios",
            back_url=url_for("options"),
        )

    @app.route("/options/reset", methods=["POST"])
    @login_required
    def reset_account():
        try:
            note_ids = [row[0] for row in db.session.query(Note.id).filter_by(user_id=current_user.id).all()]
            subject_ids = [row[0] for row in db.session.query(Subject.id).filter_by(user_id=current_user.id).all()]

            AskProfeMessage.query.filter_by(user_id=current_user.id).delete(synchronize_session=False)
            Job.query.filter_by(user_id=current_user.id).delete(synchronize_session=False)
            FlashcardDeck.query.filter_by(user_id=current_user.id).delete(synchronize_session=False)
            TaskItem.query.filter_by(user_id=current_user.id).delete(synchronize_session=False)

            if note_ids:
                NoteSourceFile.query.filter(NoteSourceFile.note_id.in_(note_ids)).delete(synchronize_session=False)
            Note.query.filter_by(user_id=current_user.id).delete(synchronize_session=False)

            if subject_ids:
                SubjectExam.query.filter(SubjectExam.subject_id.in_(subject_ids)).delete(synchronize_session=False)
            Subject.query.filter_by(user_id=current_user.id).delete(synchronize_session=False)

            StudentProfile.query.filter_by(user_id=current_user.id).delete(synchronize_session=False)

            db.session.commit()
            flash("Cuenta reseteada ✅", "success")
        except Exception:
            db.session.rollback()
            flash("Error reseteando la cuenta.", "error")
        return redirect(url_for("setup"))

    # ---------- API: fechas por asignatura ----------
    @app.route("/api/exam-dates")
    @login_required
    def api_exam_dates():
        subject_id = request.args.get("subject_id", type=int)
        if not subject_id:
            return jsonify({"dates": []})

        subject = Subject.query.filter_by(id=subject_id, user_id=current_user.id).first()
        if not subject:
            return jsonify({"dates": []})

        rows = (
            db.session.query(SubjectExam.exam_date, SubjectExam.tema)
            .filter(SubjectExam.subject_id == subject_id)
            .order_by(SubjectExam.exam_date.asc(), SubjectExam.tema.asc())
            .all()
        )
        return jsonify(
            {
                "dates": [
                    {"date": r[0].isoformat(), "tema": r[1]}
                    for r in rows
                    if r[0] is not None and r[1]
                ]
            }
        )

    # ---------- API: títulos por asignatura ----------
    @app.route("/api/titles")
    @login_required
    def api_titles():
        subject_id = request.args.get("subject_id", type=int)
        if not subject_id:
            return jsonify({"titles": []})

        subject = Subject.query.filter_by(id=subject_id, user_id=current_user.id).first()
        if not subject:
            return jsonify({"titles": []})

        rows = (
            db.session.query(Note.title)
            .filter(Note.user_id == current_user.id, Note.subject_id == subject_id)
            .distinct()
            .order_by(Note.title.asc())
            .all()
        )
        return jsonify({"titles": [r[0] for r in rows if r[0]]})

    @app.template_filter("note_fmt")
    def note_fmt_filter(text: str):
        return simple_format_note(text)

    @app.route("/api/jobs/queue")
    @login_required
    def api_jobs_queue():
        jobs = (
            Job.query.filter_by(user_id=current_user.id)
            .order_by(Job.created_at.desc())
            .limit(15)
            .all()
        )
        return jsonify(
            {
                "jobs": [
                    {
                        "id": j.id,
                        "type": j.type,
                        "status": j.status,
                        "result": j.result_message,
                        "error": j.error_message,
                        "created_at": j.created_at.isoformat(),
                        "updated_at": j.updated_at.isoformat() if j.updated_at else None,
                    }
                    for j in jobs
                ]
            }
        )

    @app.route("/api/profe/status")
    @login_required
    def api_profe_status():
        return jsonify({"profe_busy": profe_is_busy(current_user.id), "llm_busy": queue_has_work()})

    @app.route("/api/jobs/updates")
    @login_required
    def api_jobs_updates():
        try:
            jobs = (
                Job.query.filter(
                    Job.user_id == current_user.id,
                    Job.status.in_(("success", "error")),
                    Job.notified.is_(False),
                )
                .order_by(Job.updated_at.desc())
                .all()
            )
            payload = [
                {
                    "id": j.id,
                    "status": j.status,
                    "message": j.result_message or j.error_message or "",
                }
                for j in jobs
            ]
            for j in jobs:
                j.notified = True
            db.session.commit()
        except OperationalError as exc:
            db.session.rollback()
            if "database is locked" in str(exc).lower():
                return jsonify({"jobs": [], "llm_busy": queue_has_work()})
            raise
        return jsonify({"jobs": payload, "llm_busy": queue_has_work()})

    # ---------- Ask Profe ----------
    @app.route("/ask-profe", methods=["GET", "POST"])
    @login_required
    @limiter.limit("20 per minute", methods=["POST"])
    def ask_profe():
        now = datetime.utcnow()
        lock = get_active_profe_lock(now=now)
        if lock and lock.user_id != current_user.id:
            flash("El profe está atendiendo a otro estudiante. Vuelve en unos minutos.", "error")
            return redirect(url_for("dashboard"))

        if not lock:
            if request.method == "POST":
                flash("Tu tiempo con el profe ha terminado. Vuelve al menú para iniciar otro turno.", "error")
                return redirect(url_for("dashboard"))
            if queue_has_work():
                flash("El profe está ocupado procesando trabajos. Cuando termine, podrás usar «Pregúntale al profe».", "error")
                return redirect(url_for("dashboard"))
            lock = start_profe_lock(current_user.id, now=now)

        lock_ends_at = lock.ends_at if lock else now
        if lock_ends_at <= now:
            flash("Tu tiempo con el profe ha terminado.", "error")
            return redirect(url_for("dashboard"))

        available_models = fetch_models(app)
        selected_model = resolve_default_model(app, current_user.id, available_models)
        messages = fetch_ask_profe_history(current_user.id)
        profile = StudentProfile.query.filter_by(user_id=current_user.id).first()
        student_context = ""
        if profile:
            student_context = (
                f"Estudiante: {profile.student_name} (edad {profile.age}). "
                f"Características: {profile.personality_notes}"
            )

        if request.method == "POST":
            if datetime.utcnow() >= lock_ends_at:
                flash("Tu tiempo con el profe ha terminado.", "error")
                return redirect(url_for("dashboard"))
            selected_model = request.form.get("model") or selected_model
            question = request.form.get("question", "").strip()

            if question:
                try:
                    context_messages = messages[-2:]  # ultimo turno (usuario + profe)
                    history = list(context_messages)
                    history.append({"role": "user", "content": question})

                    payload = {
                        "model": selected_model,
                        "messages": [
                            {
                                "role": "system",
                                "content": (
                                    "Eres un profesor humano paciente y amable. Responde de forma clara, completa y conversacional, "
                                    "sin usar asteriscos ni acciones roleplay; escribe como hablarías en la vida real. "
                                    "Adapta tu respuesta al perfil del estudiante, pero no menciones esos datos de forma explícita "
                                    "a menos que sea relevante para la pregunta o el estudiante lo pida. "
                                    f"{student_context}"
                                ),
                            },
                            *history,
                        ],
                        "temperature": 0.7,
                    }
                    api_base = app.config["LMSTUDIO_API_BASE"].rstrip("/")
                    resp = requests.post(f"{api_base}/chat/completions", json=payload, timeout=app.config["LMSTUDIO_TIMEOUT"])
                    resp.raise_for_status()
                    data = resp.json()
                    answer = data["choices"][0]["message"]["content"]
                    history.append({"role": "assistant", "content": answer})
                    db.session.add(AskProfeMessage(user_id=current_user.id, role="user", content=question))
                    db.session.add(AskProfeMessage(user_id=current_user.id, role="assistant", content=answer))
                    prune_ask_profe_history(current_user.id, keep=12)
                    db.session.commit()
                    messages = fetch_ask_profe_history(current_user.id)
                except Timeout:
                    flash("El modelo tardó demasiado en responder.", "error")
                    return redirect(url_for("dashboard"))
                except RequestException:
                    flash("No he podido conectar con LM Studio. ¿Está encendido?", "error")
                    return redirect(url_for("dashboard"))
                except Exception:
                    db.session.rollback()
                    flash("Error inesperado procesando tu pregunta.", "error")
                    return redirect(url_for("dashboard"))

        remaining_seconds = max(0, int(math.ceil((lock_ends_at - datetime.utcnow()).total_seconds())))
        remaining_label = f"{remaining_seconds // 60:02d}:{remaining_seconds % 60:02d}"
        lock_ends_at_ts = int(lock_ends_at.replace(tzinfo=timezone.utc).timestamp() * 1000)
        return render_template(
            "ask_profe.html",
            messages=messages,
            models=available_models,
            selected_model=selected_model,
            lock_ends_at_ts=lock_ends_at_ts,
            lock_remaining_label=remaining_label,
        )

    # ---------- Subir apuntes / resumen (igual que antes) ----------
    @app.route("/add-notes", methods=["GET", "POST"])
    @login_required
    def add_notes():
        subjects = Subject.query.filter_by(user_id=current_user.id).order_by(Subject.name.asc()).all()
        available_models = fetch_models(app)
        selected_model = resolve_default_model(app, current_user.id, available_models)
        llm_limits = get_llm_limits(current_user.id)
        chunk_tokens = llm_limits["chunk_tokens"]
        chunk_overlap = llm_limits["chunk_overlap"]

        if request.method == "POST":
            selected_model = request.form.get("model") or selected_model

            subject_choice = (request.form.get("subject_choice") or "").strip()
            new_subject_name = (request.form.get("new_subject_name") or "").strip()
            new_subject_color = (request.form.get("new_subject_color") or "").strip()
            normalized_color = normalize_subject_color(new_subject_color)

            if subject_choice == "__new__":
                if not new_subject_name:
                    flash("Escribe el nombre de la nueva asignatura.", "error")
                    return redirect(url_for("add_notes"))
                if new_subject_color and not normalized_color:
                    flash("Color de asignatura inválido.", "error")
                    return redirect(url_for("add_notes"))
                subject = Subject.query.filter_by(user_id=current_user.id, name=new_subject_name).first()
                if not subject:
                    subject = Subject(
                        user_id=current_user.id,
                        name=new_subject_name,
                        color=normalized_color,
                    )
                    db.session.add(subject)
                    db.session.commit()
                elif normalized_color and subject.color != normalized_color:
                    subject.color = normalized_color
                    db.session.commit()
            else:
                try:
                    subject_id = int(subject_choice)
                except ValueError:
                    flash("Selecciona una asignatura válida.", "error")
                    return redirect(url_for("add_notes"))
                subject = Subject.query.filter_by(id=subject_id, user_id=current_user.id).first()
                if not subject:
                    flash("Asignatura inválida.", "error")
                    return redirect(url_for("add_notes"))

            title = (request.form.get("title") or "").strip()

            exam_str = ((request.form.get("exam_date_manual") or "").strip() or (request.form.get("exam_date_choice") or "").strip())
            exam_date = None
            if exam_str:
                try:
                    exam_date = datetime.strptime(exam_str, "%Y-%m-%d").date()
                except ValueError:
                    flash("Formato de fecha inválido. Usa YYYY-MM-DD.", "error")
                    return redirect(url_for("add_notes"))

            manual_text = (request.form.get("manual_text") or "").strip()
            manual_mode = request.form.get("manual_mode") == "on"
            manual_file_mode = request.form.get("manual_file_mode") == "on"
            auto_flashcards = request.form.get("auto_flashcards") == "on"

            uploads = [u for u in request.files.getlist("file") if u and u.filename]
            file_texts: list[str] = []
            filenames: list[str] = []
            source_bytes = None
            source_mime = None
            source_filename = None
            if uploads:
                for upload in uploads:
                    if not allowed_file(upload.filename):
                        flash("Solo se permiten archivos .txt, .pdf o .pptx.", "error")
                        return redirect(url_for("add_notes"))
                    filename = secure_filename(upload.filename)
                    file_bytes = upload.read()
                    if not file_bytes:
                        flash("El archivo está vacío.", "error")
                        return redirect(url_for("add_notes"))
                    if len(file_bytes) > MAX_UPLOAD_BYTES:
                        flash(f"Archivo demasiado grande. Máximo {MAX_UPLOAD_BYTES // 1024} KB.", "error")
                        return redirect(url_for("add_notes"))
                    ext = filename.rsplit(".", 1)[1].lower()
                    if ext == "pdf":
                        file_mime = upload.mimetype or "application/pdf"
                        file_text = extract_pdf_text(file_bytes)
                    elif ext == "pptx":
                        file_mime = (
                            upload.mimetype
                            or "application/vnd.openxmlformats-officedocument.presentationml.presentation"
                        )
                        file_text = extract_pptx_text(file_bytes)
                    else:
                        file_mime = upload.mimetype or "text/plain"
                        try:
                            file_text = file_bytes.decode("utf-8")
                        except UnicodeDecodeError:
                            file_text = file_bytes.decode("utf-8", errors="ignore")
                    file_text = (file_text or "").strip()
                    if not file_text:
                        flash("El archivo no contiene texto legible.", "error")
                        return redirect(url_for("add_notes"))
                    file_texts.append(file_text)
                    filenames.append(filename)
                    if len(uploads) == 1:
                        source_bytes = file_bytes
                        source_mime = file_mime
                        source_filename = filename
                combined_text = "\n\n".join(file_texts)
                if len(uploads) > 1:
                    source_filename = "archivos_combinados.txt"
                    source_bytes = combined_text.encode("utf-8")
                    source_mime = "text/plain"
            else:
                combined_text = ""
                filename = None

            filename = source_filename if uploads else None

            content_text = manual_text if manual_text else combined_text
            if manual_mode or manual_file_mode:
                if not content_text:
                    flash("Si eliges guardar sin IA, sube un TXT/PDF/PPTX o escribe contenido.", "error")
                    return redirect(url_for("add_notes"))
                # default title: nombre de archivo sin extensión + " examen " + fecha + " creado " + fecha de subida
                if title.strip():
                    final_title = title.strip()
                else:
                    base_name = Path(filename).stem if filename else subject.name
                    created_str = datetime.utcnow().date().isoformat()
                    exam_part = exam_date.isoformat() if exam_date else ""
                    final_title = f"{base_name} examen {exam_part} creado {created_str}".strip()
                note = Note(
                    user_id=current_user.id,
                    subject_id=subject.id,
                    title=final_title,
                    exam_date=exam_date,
                    original_filename=filename,
                    content=content_text,
                    ai_used=False,
                )
                db.session.add(note)
                db.session.commit()
                flash("Apuntes guardados (sin IA) ✅", "success")
                return redirect(url_for("dashboard"))

            if not file_texts:
                flash("Para usar IA sube un TXT, PDF o PPTX válido.", "error")
                return redirect(url_for("add_notes"))

            base_name = Path(filename or "input").stem
            if title.strip():
                final_title = title.strip()
            else:
                created_str = datetime.utcnow().date().isoformat()
                exam_part = exam_date.isoformat() if exam_date else ""
                final_title = f"{base_name} examen {exam_part} creado {created_str} ({selected_model})".strip()
            exam_date_str = exam_date.isoformat() if exam_date else "No indicada"
            chunks = []
            for text in file_texts:
                chunks.extend(
                    chunk_text_with_overlap(
                        text,
                        max_tokens=chunk_tokens,
                        overlap=chunk_overlap,
                    )
                )
            if not chunks:
                flash("No se pudo dividir el texto en fragmentos para IA.", "error")
                return redirect(url_for("add_notes"))

            note = Note(
                user_id=current_user.id,
                subject_id=subject.id,
                title=final_title,
                exam_date=exam_date,
                original_filename=filename,
                content=f"{final_title}\n\n",
                ai_used=True,
            )
            db.session.add(note)
            db.session.flush()
            if source_bytes:
                db.session.add(
                    NoteSourceFile(
                        note_id=note.id,
                        filename=filename or "input.txt",
                        content_type=source_mime or "application/octet-stream",
                        data=source_bytes,
                    )
                )
            db.session.commit()

            deck = None
            if auto_flashcards:
                deck_title = None
                if exam_date:
                    exam_row = (
                        SubjectExam.query.filter_by(subject_id=subject.id, exam_date=exam_date)
                        .order_by(SubjectExam.tema.asc())
                        .first()
                    )
                    if exam_row and exam_row.tema:
                        deck_title = exam_row.tema
                if not deck_title:
                    deck_title = title.strip() if title.strip() else base_name or final_title
                deck = FlashcardDeck.query.filter_by(
                    user_id=current_user.id,
                    subject_id=subject.id,
                    exam_date=exam_date,
                    title=deck_title,
                ).first()
                if not deck:
                    deck = FlashcardDeck(
                        user_id=current_user.id,
                        subject_id=subject.id,
                        title=deck_title,
                        exam_date=exam_date,
                        source_note_id=note.id,
                        flashcards=[],
                    )
                    db.session.add(deck)
                    db.session.flush()

            deck_size_before = len(deck.flashcards or []) if deck else 0
            for idx, chunk in enumerate(chunks):
                job = Job(
                    user_id=current_user.id,
                    type="note_ai_chunk",
                    payload={
                        "user_id": current_user.id,
                        "subject_id": subject.id,
                        "note_id": note.id,
                        "title": final_title,
                        "exam_date": exam_date_str,
                        "filename": filename or "input.txt",
                        "text": chunk,
                        "model": selected_model,
                        "chunk_index": idx,
                        "total_chunks": len(chunks),
                    },
                )
                db.session.add(job)
                if auto_flashcards and deck:
                    db.session.add(
                        Job(
                            user_id=current_user.id,
                            type="flashcards_ai_chunk",
                            payload={
                                "user_id": current_user.id,
                                "note_id": note.id,
                                "deck_id": deck.id,
                                "deck_size_before": deck_size_before,
                                "model": selected_model,
                                "count": FLASHCARD_CHUNK_DEFAULT,
                                "chunk_index": idx,
                                "total_chunks": len(chunks),
                                "text": chunk,
                            },
                        )
                    )
            db.session.commit()

            if auto_flashcards and deck:
                total_cards = len(chunks) * FLASHCARD_CHUNK_DEFAULT
                flash(
                    "Resumen encolado en "
                    f"{len(chunks)} fragmento(s) ✅ Flashcards en cola: "
                    f"{len(chunks)} × {FLASHCARD_CHUNK_DEFAULT} (total estimado {total_cards}).",
                    "success",
                )
            else:
                flash(
                    f"Resumen encolado en {len(chunks)} fragmento(s) ✅ Se irá completando a medida que procesamos cada parte.",
                    "success",
                )
            return redirect(url_for("add_notes"))

        jobs = (
            Job.query.filter_by(user_id=current_user.id)
            .order_by(Job.created_at.desc())
            .limit(10)
            .all()
        )

        return render_template(
            "add_notes.html",
            subjects=subjects,
            models=available_models,
            selected_model=selected_model,
            max_kb=MAX_UPLOAD_BYTES // 1024,
            chunk_tokens=chunk_tokens,
            chunk_overlap=chunk_overlap,
            jobs=jobs,
        )

    # ---------- Consultar apuntes ----------
    @app.route("/notes")
    @login_required
    def notes_list():
        subjects = Subject.query.filter_by(user_id=current_user.id).order_by(Subject.name.asc()).all()
        subject_id = request.args.get("subject_id", type=int)
        tema_q = (request.args.get("tema") or "").strip()
        title_q = (request.args.get("title") or "").strip()

        q = Note.query.filter_by(user_id=current_user.id)
        if subject_id:
            q = q.filter(Note.subject_id == subject_id)
        if tema_q:
            q = q.join(
                SubjectExam,
                (SubjectExam.subject_id == Note.subject_id) & (SubjectExam.exam_date == Note.exam_date),
            ).filter(SubjectExam.tema == tema_q)
        if title_q:
            q = q.filter(Note.title.ilike(f"%{title_q}%"))

        notes = q.order_by(Note.updated_at.desc()).all()
        tema_options = fetch_tema_options(current_user.id, subject_id)
        filters_active = bool(subject_id or tema_q or title_q)
        chunk_totals: dict[int, int] = {}
        note_models: dict[int, str] = {}
        source_files: dict[int, NoteSourceFile] = {}
        note_temas: dict[int, list[str]] = {}
        if notes:
            note_ids = {n.id for n in notes}
            exam_pairs = {(n.subject_id, n.exam_date) for n in notes if n.exam_date}
            jobs = Job.query.filter(
                Job.user_id == current_user.id,
                Job.type == "note_ai_chunk",
            ).all()
            for job in jobs:
                payload = job.payload or {}
                note_id = payload.get("note_id")
                if note_id not in note_ids:
                    continue
                try:
                    total_chunks = int(payload.get("total_chunks") or 0)
                except (ValueError, TypeError):
                    total_chunks = 0
                if total_chunks:
                    prev = chunk_totals.get(note_id, 0)
                    chunk_totals[note_id] = max(prev, total_chunks)
                model = payload.get("model")
                if model:
                    note_models.setdefault(note_id, model)
            source_rows = NoteSourceFile.query.filter(NoteSourceFile.note_id.in_(note_ids)).all()
            source_files = {s.note_id: s for s in source_rows}
            if exam_pairs:
                subject_ids = {sid for sid, _ in exam_pairs}
                exam_dates = {d for _, d in exam_pairs}
                tema_rows = SubjectExam.query.filter(
                    SubjectExam.subject_id.in_(subject_ids),
                    SubjectExam.exam_date.in_(exam_dates),
                ).all()
                temas_by_pair: dict[tuple, list[str]] = {}
                for row in tema_rows:
                    key = (row.subject_id, row.exam_date)
                    temas_by_pair.setdefault(key, []).append(row.tema)
                for note in notes:
                    if not note.exam_date:
                        continue
                    temas = temas_by_pair.get((note.subject_id, note.exam_date))
                    if temas:
                        note_temas[note.id] = sorted(set(temas))

        return render_template(
            "notes_list.html",
            subjects=subjects,
            notes=notes,
            filters={"subject_id": subject_id or "", "tema": tema_q, "title": title_q},
            chunk_totals=chunk_totals,
            note_models=note_models,
            source_files=source_files,
            note_temas=note_temas,
            temas=tema_options,
            filters_active=filters_active,
        )

    @app.route("/notes/<int:note_id>/source-file")
    @login_required
    def note_source_file(note_id: int):
        note = Note.query.filter_by(id=note_id, user_id=current_user.id).first_or_404()
        source = NoteSourceFile.query.filter_by(note_id=note.id).first()
        if not source:
            flash("No se encontró el archivo original.", "error")
            return redirect(url_for("notes_list"))
        return send_file(
            BytesIO(source.data),
            mimetype=source.content_type or "application/octet-stream",
            download_name=source.filename or "archivo_original",
            as_attachment=True,
        )

    @app.route("/notes/merge", methods=["POST"])
    @login_required
    def notes_merge():
        target_id = request.form.get("target_note", type=int)
        source_id = request.form.get("source_note", type=int)
        if not target_id or not source_id or target_id == source_id:
            flash("Elige apuntes válidos (origen y destino distintos).", "error")
            return redirect(url_for("notes_list"))

        target = Note.query.filter_by(id=target_id, user_id=current_user.id).first()
        source = Note.query.filter_by(id=source_id, user_id=current_user.id).first()
        if not target or not source:
            flash("No se encontraron los apuntes seleccionados.", "error")
            return redirect(url_for("notes_list"))

        merged = (target.content or "") + "\n\n" + (source.content or "")
        target.content = merged.strip()
        target.updated_at = datetime.utcnow()
        db.session.commit()
        flash("Apuntes combinados ✅", "success")
        return redirect(url_for("notes_list"))

    @app.route("/notes/<int:note_id>/edit", methods=["GET", "POST"])
    @login_required
    def note_edit(note_id: int):
        note = Note.query.filter_by(id=note_id, user_id=current_user.id).first_or_404()
        subjects = Subject.query.filter_by(user_id=current_user.id).order_by(Subject.name.asc()).all()

        if request.method == "POST":
            subject_id = request.form.get("subject_id", type=int)
            subject = Subject.query.filter_by(id=subject_id, user_id=current_user.id).first()
            if not subject:
                flash("Asignatura inválida.", "error")
                return redirect(url_for("note_edit", note_id=note.id))

            title = (request.form.get("title") or "").strip()
            if not title:
                # si no hay título, usamos la primera línea del contenido
                title = (request.form.get("content") or "").strip().splitlines()[0] if (request.form.get("content") or "").strip() else ""
            if not title:
                flash("El título es obligatorio.", "error")
                return redirect(url_for("note_edit", note_id=note.id))

            exam_date_str = (request.form.get("exam_date") or "").strip()
            exam_date = None
            if exam_date_str:
                try:
                    exam_date = datetime.strptime(exam_date_str, "%Y-%m-%d").date()
                except ValueError:
                    flash("Formato de fecha inválido (YYYY-MM-DD).", "error")
                    return redirect(url_for("note_edit", note_id=note.id))

            content = (request.form.get("content") or "").strip()
            if not content:
                flash("El contenido no puede estar vacío.", "error")
                return redirect(url_for("note_edit", note_id=note.id))

            note.subject_id = subject.id
            note.title = title
            note.exam_date = exam_date
            note.content = content
            note.ai_used = False
            note.original_filename = None
            NoteSourceFile.query.filter_by(note_id=note.id).delete()

            db.session.commit()
            flash("Apunte actualizado ✅", "success")
            return redirect(url_for("notes_list"))

        return render_template("note_edit.html", note=note, subjects=subjects)

    @app.route("/notes/<int:note_id>/delete", methods=["POST"])
    @login_required
    def note_delete(note_id: int):
        note = Note.query.filter_by(id=note_id, user_id=current_user.id).first_or_404()
        if has_incomplete_jobs_for_payload(current_user.id, "note_id", note.id):
            flash("No puedes borrar este apunte mientras haya trabajos en curso.", "error")
            return redirect(url_for("notes_list"))
        NoteSourceFile.query.filter_by(note_id=note.id).delete()
        db.session.delete(note)
        db.session.commit()
        flash("Apunte borrado ✅", "success")
        return redirect(url_for("notes_list"))

    # ===========================
    # FLASHCARDS
    # ===========================

    @app.route("/flashcards/create", methods=["GET", "POST"])
    @login_required
    def flashcards_create():
        subjects = Subject.query.filter_by(user_id=current_user.id).order_by(Subject.name.asc()).all()
        notes = Note.query.filter_by(user_id=current_user.id).order_by(Note.updated_at.desc()).all()
        decks = (
            FlashcardDeck.query.filter_by(user_id=current_user.id)
            .order_by(FlashcardDeck.updated_at.desc())
            .all()
        )
        available_models = fetch_models(app)
        selected_model = resolve_default_model(app, current_user.id, available_models)
        llm_limits = get_llm_limits(current_user.id)
        chunk_tokens = llm_limits["chunk_tokens"]
        chunk_overlap = llm_limits["chunk_overlap"]
        note_chunk_map = build_note_chunks_map(
            current_user.id,
            notes,
            max_tokens=chunk_tokens,
            overlap=chunk_overlap,
        )

        if request.method == "POST":
            mode = (request.form.get("mode") or "ai").strip()
            selected_model = request.form.get("model") or selected_model

            if mode == "ai":
                deck_id = request.form.get("deck_id", type=int)
                target_deck = None
                if deck_id:
                    target_deck = FlashcardDeck.query.filter_by(id=deck_id, user_id=current_user.id).first()
                    if not target_deck:
                        flash("Deck de destino inválido.", "error")
                        return redirect(url_for("flashcards_create"))

                note_id = request.form.get("note_id", type=int)
                if not note_id:
                    flash("Selecciona un resumen/apunte para generar flashcards.", "error")
                    return redirect(url_for("flashcards_create"))

                note = Note.query.filter_by(id=note_id, user_id=current_user.id).first()
                if not note:
                    flash("Resumen/apunte inválido.", "error")
                    return redirect(url_for("flashcards_create"))

                custom_title = (request.form.get("ai_deck_title") or "").strip()
                count_per_chunk = request.form.get("count", type=int) or FLASHCARD_CHUNK_DEFAULT
                if count_per_chunk not in FLASHCARD_CHUNK_COUNTS:
                    flash("Selecciona una cantidad válida de flashcards por fragmento.", "error")
                    return redirect(url_for("flashcards_create"))

                chunk_entry = note_chunk_map.get(note.id) or {"chunks": [note.content or ""], "total": 1}
                chunks = chunk_entry.get("chunks") or [note.content or ""]
                chunk_count = len(chunks)

                if target_deck:
                    deck = target_deck
                else:
                    deck = FlashcardDeck(
                        user_id=current_user.id,
                        subject_id=note.subject_id,
                        title=custom_title or note.title,
                        exam_date=note.exam_date,
                        source_note_id=note.id,
                        flashcards=[],
                    )
                    db.session.add(deck)
                    db.session.commit()

                deck_size_before = len(deck.flashcards or [])
                for idx, chunk in enumerate(chunks):
                    job = Job(
                        user_id=current_user.id,
                        type="flashcards_ai_chunk",
                        payload={
                            "user_id": current_user.id,
                            "note_id": note.id,
                            "deck_id": deck.id,
                            "deck_size_before": deck_size_before,
                            "model": selected_model,
                            "count": count_per_chunk,
                            "chunk_index": idx,
                            "total_chunks": chunk_count,
                            "text": chunk,
                        },
                    )
                    db.session.add(job)
                db.session.commit()
                total_cards = count_per_chunk * chunk_count
                flash(f"Generación encolada ✅ {chunk_count} fragmentos × {count_per_chunk} (total estimado {total_cards}).", "success")
                return redirect(url_for("flashcards_create"))

            # ---- MANUAL ----
            deck_id = request.form.get("deck_id", type=int)
            target_deck = None
            if deck_id:
                target_deck = FlashcardDeck.query.filter_by(id=deck_id, user_id=current_user.id).first()
                if not target_deck:
                    flash("Deck de destino inválido.", "error")
                    return redirect(url_for("flashcards_create"))

            subject_choice = (request.form.get("subject_choice") or "").strip()
            new_subject_name = (request.form.get("new_subject_name") or "").strip()
            new_subject_color = (request.form.get("new_subject_color") or "").strip()
            normalized_color = normalize_subject_color(new_subject_color)

            if not target_deck:
                if subject_choice == "__new__":
                    if not new_subject_name:
                        flash("Escribe el nombre de la nueva asignatura.", "error")
                        return redirect(url_for("flashcards_create"))
                    if new_subject_color and not normalized_color:
                        flash("Color de asignatura inválido.", "error")
                        return redirect(url_for("flashcards_create"))
                    subject = Subject.query.filter_by(user_id=current_user.id, name=new_subject_name).first()
                    if not subject:
                        subject = Subject(
                            user_id=current_user.id,
                            name=new_subject_name,
                            color=normalized_color,
                        )
                        db.session.add(subject)
                        db.session.commit()
                    elif normalized_color and subject.color != normalized_color:
                        subject.color = normalized_color
                        db.session.commit()
                else:
                    try:
                        subject_id = int(subject_choice)
                    except ValueError:
                        flash("Selecciona una asignatura válida.", "error")
                        return redirect(url_for("flashcards_create"))
                    subject = Subject.query.filter_by(id=subject_id, user_id=current_user.id).first()
                    if not subject:
                        flash("Asignatura inválida.", "error")
                        return redirect(url_for("flashcards_create"))

                title = (request.form.get("title") or "").strip()
                if not title:
                    flash("El título es obligatorio.", "error")
                    return redirect(url_for("flashcards_create"))

                exam_str = (request.form.get("exam_date") or "").strip()
                exam_date = None
                if exam_str:
                    try:
                        exam_date = datetime.strptime(exam_str, "%Y-%m-%d").date()
                    except ValueError:
                        flash("Formato de fecha inválido (YYYY-MM-DD).", "error")
                        return redirect(url_for("flashcards_create"))

            q = (request.form.get("q") or "").strip()
            a = (request.form.get("a") or "").strip()
            b = (request.form.get("b") or "").strip()
            c = (request.form.get("c") or "").strip()
            d = (request.form.get("d") or "").strip()
            correct = request.form.get("correct", type=int)

            if not q or not a or not b or not c or not d or correct is None:
                flash("Rellena pregunta, 4 respuestas y marca la correcta.", "error")
                return redirect(url_for("flashcards_create"))

            if correct not in (0, 1, 2, 3):
                flash("Índice de respuesta correcta inválido.", "error")
                return redirect(url_for("flashcards_create"))

            cards = [{"question": q, "options": [a, b, c, d], "correct_index": correct}]

            if target_deck:
                target_deck.flashcards = (target_deck.flashcards or []) + cards
                db.session.commit()
                flash("Flashcard añadida al deck existente ✅", "success")
                return redirect(url_for("flashcards_list"))
            else:
                deck = FlashcardDeck(
                    user_id=current_user.id,
                    subject_id=subject.id,
                    title=title,
                    exam_date=exam_date,
                    source_note_id=None,
                    flashcards=cards,
                )
                db.session.add(deck)
                db.session.commit()
                flash("Flashcard creada y guardada ✅", "success")
                return redirect(url_for("flashcards_list"))

        return render_template(
            "flashcards_create.html",
            subjects=subjects,
            notes=notes,
            note_chunk_counts={nid: info.get("total", 1) for nid, info in note_chunk_map.items()},
            models=available_models,
            selected_model=selected_model,
            chunk_tokens=chunk_tokens,
            chunk_overlap=chunk_overlap,
            decks=decks,
            count_options=FLASHCARD_CHUNK_COUNTS,
            default_count=FLASHCARD_CHUNK_DEFAULT,
            jobs=Job.query.filter_by(user_id=current_user.id).order_by(Job.created_at.desc()).limit(10).all(),
        )

    @app.route("/flashcards")
    @login_required
    def flashcards_list():
        subjects = Subject.query.filter_by(user_id=current_user.id).order_by(Subject.name.asc()).all()
        subject_id = request.args.get("subject_id", type=int)
        tema_q = (request.args.get("tema") or "").strip()
        title_q = (request.args.get("title") or "").strip()

        q = FlashcardDeck.query.filter_by(user_id=current_user.id)

        if subject_id:
            q = q.filter(FlashcardDeck.subject_id == subject_id)

        if tema_q:
            q = q.join(
                SubjectExam,
                (SubjectExam.subject_id == FlashcardDeck.subject_id)
                & (SubjectExam.exam_date == FlashcardDeck.exam_date),
            ).filter(SubjectExam.tema == tema_q)

        if title_q:
            q = q.filter(FlashcardDeck.title.ilike(f"%{title_q}%"))

        decks = q.order_by(FlashcardDeck.updated_at.desc()).all()
        deck_temas: dict[int, list[str]] = {}
        if decks:
            exam_pairs = {(d.subject_id, d.exam_date) for d in decks if d.exam_date}
            if exam_pairs:
                subject_ids = {sid for sid, _ in exam_pairs}
                exam_dates = {d for _, d in exam_pairs}
                tema_rows = SubjectExam.query.filter(
                    SubjectExam.subject_id.in_(subject_ids),
                    SubjectExam.exam_date.in_(exam_dates),
                ).all()
                temas_by_pair: dict[tuple, list[str]] = {}
                for row in tema_rows:
                    key = (row.subject_id, row.exam_date)
                    temas_by_pair.setdefault(key, []).append(row.tema)
                for deck in decks:
                    if not deck.exam_date:
                        continue
                    temas = temas_by_pair.get((deck.subject_id, deck.exam_date))
                    if temas:
                        deck_temas[deck.id] = sorted(set(temas))
        tema_options = fetch_tema_options(current_user.id, subject_id)
        filters_active = bool(subject_id or tema_q or title_q)

        return render_template(
            "flashcards_list.html",
            subjects=subjects,
            decks=decks,
            filters={"subject_id": subject_id or "", "tema": tema_q, "title": title_q},
            temas=tema_options,
            filters_active=filters_active,
            deck_temas=deck_temas,
        )

    @app.route("/flashcards/<int:deck_id>/edit", methods=["GET", "POST"])
    @login_required
    def flashcards_edit(deck_id: int):
        deck = FlashcardDeck.query.filter_by(id=deck_id, user_id=current_user.id).first_or_404()
        subjects = Subject.query.filter_by(user_id=current_user.id).order_by(Subject.name.asc()).all()
        notes = Note.query.filter_by(user_id=current_user.id).order_by(Note.updated_at.desc()).all()
        other_decks = (
            FlashcardDeck.query.filter(FlashcardDeck.user_id == current_user.id, FlashcardDeck.id != deck.id)
            .order_by(FlashcardDeck.updated_at.desc())
            .all()
        )
        available_models = fetch_models(app)
        selected_model = resolve_default_model(app, current_user.id, available_models)
        llm_limits = get_llm_limits(current_user.id)
        chunk_tokens = llm_limits["chunk_tokens"]
        chunk_overlap = llm_limits["chunk_overlap"]
        note_chunk_map = build_note_chunks_map(
            current_user.id,
            notes,
            max_tokens=chunk_tokens,
            overlap=chunk_overlap,
        )

        if request.method == "POST":
            mode = (request.form.get("mode") or "manual").strip()

            if mode == "append_ai":
                note_id = request.form.get("note_id", type=int)
                model = request.form.get("model") or selected_model
                if not note_id:
                    flash("Selecciona un resumen/apunte para generar flashcards.", "error")
                    return redirect(url_for("flashcards_edit", deck_id=deck.id))

                note = Note.query.filter_by(id=note_id, user_id=current_user.id).first()
                if not note:
                    flash("Resumen/apunte inválido.", "error")
                    return redirect(url_for("flashcards_edit", deck_id=deck.id))

                count_per_chunk = request.form.get("count", type=int) or FLASHCARD_CHUNK_DEFAULT
                if count_per_chunk not in FLASHCARD_CHUNK_COUNTS:
                    flash("Selecciona una cantidad válida de flashcards por fragmento.", "error")
                    return redirect(url_for("flashcards_edit", deck_id=deck.id))
                chunk_entry = note_chunk_map.get(note.id) or {"chunks": [note.content or ""], "total": 1}
                chunks = chunk_entry.get("chunks") or [note.content or ""]
                chunk_count = len(chunks)

                deck_size_before = len(deck.flashcards or [])
                for idx, chunk in enumerate(chunks):
                    job = Job(
                        user_id=current_user.id,
                        type="flashcards_ai_chunk",
                        payload={
                            "user_id": current_user.id,
                            "note_id": note.id,
                            "deck_id": deck.id,
                            "deck_size_before": deck_size_before,
                            "model": model,
                            "count": count_per_chunk,
                            "chunk_index": idx,
                            "total_chunks": chunk_count,
                            "text": chunk,
                        },
                    )
                    db.session.add(job)
                db.session.commit()
                total_cards = count_per_chunk * chunk_count
                flash(f"Generación de flashcards encolada ✅ {chunk_count} fragmentos × {count_per_chunk} (total estimado {total_cards}).", "success")
                return redirect(url_for("flashcards_edit", deck_id=deck.id))

            if mode == "merge":
                merge_deck_id = request.form.get("merge_deck_id", type=int)
                source = None
                if merge_deck_id:
                    source = FlashcardDeck.query.filter_by(id=merge_deck_id, user_id=current_user.id).first()
                if not source:
                    flash("Deck a combinar inválido.", "error")
                    return redirect(url_for("flashcards_edit", deck_id=deck.id))

                deck.flashcards = (deck.flashcards or []) + (source.flashcards or [])
                db.session.commit()
                flash(f"Decks combinados ✅ Ahora hay {len(deck.flashcards)} flashcards.", "success")
                return redirect(url_for("flashcards_edit", deck_id=deck.id))

            subject_id = request.form.get("subject_id", type=int)
            subject = Subject.query.filter_by(id=subject_id, user_id=current_user.id).first()
            if not subject:
                flash("Asignatura inválida.", "error")
                return redirect(url_for("flashcards_edit", deck_id=deck.id))

            title = (request.form.get("title") or "").strip()
            if not title:
                flash("El título es obligatorio.", "error")
                return redirect(url_for("flashcards_edit", deck_id=deck.id))

            exam_str = (request.form.get("exam_date") or "").strip()
            exam_date = None
            if exam_str:
                try:
                    exam_date = datetime.strptime(exam_str, "%Y-%m-%d").date()
                except ValueError:
                    flash("Formato de fecha inválido (YYYY-MM-DD).", "error")
                    return redirect(url_for("flashcards_edit", deck_id=deck.id))

            # Recoger cards desde form (permitiendo huecos en índices)
            cards = []
            idxs = sorted(
                {
                    int(k.split("_")[-1])
                    for k in request.form.keys()
                    if k.startswith("card_q_") and k.split("_")[-1].isdigit()
                }
            )
            for i in idxs:
                qtext = (request.form.get(f"card_q_{i}") or "").strip()
                o0 = (request.form.get(f"card_o0_{i}") or "").strip()
                o1 = (request.form.get(f"card_o1_{i}") or "").strip()
                o2 = (request.form.get(f"card_o2_{i}") or "").strip()
                o3 = (request.form.get(f"card_o3_{i}") or "").strip()
                correct = request.form.get(f"card_correct_{i}", type=int)

                # si la fila está vacía, la ignoramos
                if not qtext and not o0 and not o1 and not o2 and not o3:
                    continue

                if not qtext or not o0 or not o1 or not o2 or not o3 or correct is None:
                    flash(f"Flashcard #{i+1}: faltan campos.", "error")
                    return redirect(url_for("flashcards_edit", deck_id=deck.id))
                if correct not in (0, 1, 2, 3):
                    flash(f"Flashcard #{i+1}: correcta inválida.", "error")
                    return redirect(url_for("flashcards_edit", deck_id=deck.id))

                cards.append({"question": qtext, "options": [o0, o1, o2, o3], "correct_index": correct})

            if not cards:
                flash("Debes tener al menos 1 flashcard.", "error")
                return redirect(url_for("flashcards_edit", deck_id=deck.id))

            deck.subject_id = subject.id
            deck.title = title
            deck.exam_date = exam_date
            deck.flashcards = cards
            db.session.commit()

            flash("Flashcards actualizadas ✅", "success")
            return redirect(url_for("flashcards_list"))

        return render_template(
            "flashcards_edit.html",
            deck=deck,
            subjects=subjects,
            notes=notes,
            note_chunk_counts={nid: info.get("total", 1) for nid, info in note_chunk_map.items()},
            other_decks=other_decks,
            models=available_models,
            selected_model=selected_model,
            count_options=FLASHCARD_CHUNK_COUNTS,
            default_count=FLASHCARD_CHUNK_DEFAULT,
            jobs=Job.query.filter_by(user_id=current_user.id).order_by(Job.created_at.desc()).limit(10).all(),
        )

    @app.route("/flashcards/<int:deck_id>/delete", methods=["POST"])
    @login_required
    def flashcards_delete(deck_id: int):
        deck = FlashcardDeck.query.filter_by(id=deck_id, user_id=current_user.id).first_or_404()
        if has_incomplete_jobs_for_payload(current_user.id, "deck_id", deck.id):
            flash("No puedes borrar este deck mientras haya trabajos en curso.", "error")
            return redirect(url_for("flashcards_list"))
        db.session.delete(deck)
        db.session.commit()
        flash("Deck de flashcards borrado ✅", "success")
        return redirect(url_for("flashcards_list"))

    return app


app = create_app()
