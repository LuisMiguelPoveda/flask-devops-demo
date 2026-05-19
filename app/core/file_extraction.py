import os
from io import BytesIO

from PyPDF2 import PdfReader
from pptx import Presentation

ALLOWED_EXTENSIONS = {"txt", "pdf", "pptx"}
MAX_UPLOAD_BYTES = 50 * 1024 * 1024


def allowed_file(filename: str) -> bool:
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def save_upload_stream(upload, dest_path: str, max_bytes: int) -> tuple[int, str | None]:
    total = 0
    try:
        with open(dest_path, "wb") as handle:
            while True:
                chunk = upload.stream.read(1024 * 1024)
                if not chunk:
                    break
                total += len(chunk)
                if total > max_bytes:
                    break
                handle.write(chunk)
    except Exception:
        if os.path.exists(dest_path):
            os.remove(dest_path)
        raise

    if total == 0:
        if os.path.exists(dest_path):
            os.remove(dest_path)
        return 0, "empty"
    if total > max_bytes:
        if os.path.exists(dest_path):
            os.remove(dest_path)
        return total, "too_large"
    return total, None


def extract_pdf_text(file_bytes: bytes) -> str:
    reader = PdfReader(BytesIO(file_bytes))
    chunks: list[str] = []
    for page in reader.pages:
        try:
            txt = page.extract_text() or ""
        except Exception:
            txt = ""
        if txt:
            chunks.append(txt)
    return "\n".join(chunks)


def extract_pptx_text(file_bytes: bytes) -> str:
    prs = Presentation(BytesIO(file_bytes))
    chunks: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text or ""
                if text.strip():
                    chunks.append(text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    row_text = " ".join(cell.text.strip() for cell in row.cells if cell.text)
                    if row_text:
                        chunks.append(row_text)
    return "\n".join(chunks)
