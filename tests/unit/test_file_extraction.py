import io
import os

import pytest

from app.core.file_extraction import (
    MAX_UPLOAD_BYTES,
    allowed_file,
    extract_pdf_text,
    extract_pptx_text,
    save_upload_stream,
)


class TestAllowedFile:
    def test_pdf_allowed(self):
        assert allowed_file("notes.pdf") is True

    def test_txt_allowed(self):
        assert allowed_file("notes.txt") is True

    def test_pptx_allowed(self):
        assert allowed_file("slides.pptx") is True

    def test_exe_not_allowed(self):
        assert allowed_file("malware.exe") is False

    def test_no_extension_not_allowed(self):
        assert allowed_file("noextension") is False

    def test_dot_only_allowed(self):
        # ".pdf".rsplit(".", 1)[1] == "pdf" — implementation allows dotfiles
        assert allowed_file(".pdf") is True

    def test_case_insensitive(self):
        assert allowed_file("NOTES.PDF") is True
        assert allowed_file("Slides.PPTX") is True

    def test_double_extension_uses_last(self):
        assert allowed_file("file.exe.pdf") is True
        assert allowed_file("file.pdf.exe") is False


class FakeUpload:
    def __init__(self, data: bytes):
        self.stream = io.BytesIO(data)


class TestSaveUploadStream:
    def test_saves_file_successfully(self, tmp_path):
        data = b"hello world"
        dest = str(tmp_path / "out.txt")
        total, error = save_upload_stream(FakeUpload(data), dest, MAX_UPLOAD_BYTES)
        assert error is None
        assert total == len(data)
        assert open(dest, "rb").read() == data

    def test_empty_upload_returns_empty_error(self, tmp_path):
        dest = str(tmp_path / "out.txt")
        total, error = save_upload_stream(FakeUpload(b""), dest, MAX_UPLOAD_BYTES)
        assert error == "empty"
        assert total == 0
        assert not os.path.exists(dest)

    def test_oversized_upload_returns_too_large(self, tmp_path):
        data = b"x" * 200
        dest = str(tmp_path / "out.txt")
        total, error = save_upload_stream(FakeUpload(data), dest, max_bytes=100)
        assert error == "too_large"
        assert not os.path.exists(dest)

    def test_exact_limit_accepted(self, tmp_path):
        data = b"x" * 100
        dest = str(tmp_path / "out.txt")
        total, error = save_upload_stream(FakeUpload(data), dest, max_bytes=100)
        assert error is None


class TestExtractPdfText:
    def test_returns_string(self):
        try:
            from reportlab.pdfgen import canvas as rl_canvas

            buf = io.BytesIO()
            c = rl_canvas.Canvas(buf)
            c.drawString(100, 750, "Hello PDF")
            c.save()
            result = extract_pdf_text(buf.getvalue())
            assert isinstance(result, str)
        except ImportError:
            pytest.skip("reportlab not installed — skipping PDF content test")

    def test_empty_pdf_returns_string(self):
        try:
            from reportlab.pdfgen import canvas as rl_canvas

            buf = io.BytesIO()
            c = rl_canvas.Canvas(buf)
            c.save()
            result = extract_pdf_text(buf.getvalue())
            assert isinstance(result, str)
        except ImportError:
            pytest.skip("reportlab not installed")


class TestExtractPptxText:
    def test_extracts_text_from_slide(self):
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        slide_layout = prs.slide_layouts[5]
        slide = prs.slides.add_slide(slide_layout)
        txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        txBox.text_frame.text = "Hello PPTX"

        buf = io.BytesIO()
        prs.save(buf)
        result = extract_pptx_text(buf.getvalue())
        assert "Hello PPTX" in result

    def test_empty_presentation_returns_empty_string(self):
        from pptx import Presentation

        prs = Presentation()
        buf = io.BytesIO()
        prs.save(buf)
        result = extract_pptx_text(buf.getvalue())
        assert result == ""

    def test_extracts_table_text(self):
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(2)).table
        table.cell(0, 0).text = "Col A"
        table.cell(0, 1).text = "Col B"
        table.cell(1, 0).text = "Val 1"
        table.cell(1, 1).text = "Val 2"

        buf = io.BytesIO()
        prs.save(buf)
        result = extract_pptx_text(buf.getvalue())
        assert "Col A" in result
        assert "Val 2" in result


class TestSaveUploadStreamExceptions:
    def test_io_error_during_write_cleans_up(self, tmp_path):
        dest = str(tmp_path / "out.txt")

        class BrokenStream:
            def read(self, n):
                raise OSError("disk full")

        class BrokenUpload:
            stream = BrokenStream()

        with pytest.raises(OSError):
            save_upload_stream(BrokenUpload(), dest, MAX_UPLOAD_BYTES)
        assert not os.path.exists(dest)
